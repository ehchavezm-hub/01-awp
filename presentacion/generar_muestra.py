"""Genera presentacion/muestra.pptx: 8 láminas de muestra sobre la plantilla del repositorio.

Uso:  python presentacion/generar_muestra.py
Requiere: python-pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

RAIZ = Path(__file__).resolve().parent.parent
PLANTILLA = RAIZ / "plantilla_ppt.pptx"
SALIDA = RAIZ / "presentacion" / "muestra.pptx"
ICONOS = RAIZ / "presentacion" / "assets" / "iconos"

# Paleta AECOM de plantilla_ppt.pptx (lámina 10, "Configuración maestra").
NEGRO = "1A1A1A"      # fondo portada / oscuro
FONDO = "F2F2F2"      # fondo láminas de contenido
NARANJA = "FF6600"    # acento principal
NARANJA_OSC = "CC4400"
NARANJA_CLARO = "FF9955"  # KPIs / destacados
TEXTO = "2E2E2E"      # texto cuerpo
AZUL = "003087"       # ingeniería / técnico
AZUL_MEDIO = "0078C1"  # datos
AZUL_PALIDO = "A8CFE8"
GRIS = "6D6E71"       # texto secundario
GRIS_MEDIO = "555555"
TARJETA = "222222"    # tarjetas sobre oscuro
BORDE = "BCBEC0"
BLANCO = "FFFFFF"
FUENTE = "Arial"
MONO = "Courier New"

ANCHO, ALTO = 13.333, 7.5
MARGEN = 0.6


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- utilidades

def fondo(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def caja(slide, x, y, w, h, relleno=None, borde=None, grosor=0.75, forma=MSO_SHAPE.RECTANGLE, radio=None):
    shp = slide.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    if relleno:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(relleno)
    else:
        shp.fill.background()
    if borde:
        shp.line.color.rgb = rgb(borde)
        shp.line.width = Pt(grosor)
    else:
        shp.line.fill.background()
    if radio is not None and forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radio
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def texto(slide, x, y, w, h, parrafos, tam=14, color=TEXTO, negrita=False, alinear=PP_ALIGN.LEFT,
          ancla=MSO_ANCHOR.TOP, fuente=FUENTE, espaciado=0, interlineado=None, espacio_despues=0, nombre=None):
    """parrafos: str | lista de str | lista de listas de runs (dict: t, tam, color, negrita, fuente, esp)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if nombre:
        tb.name = nombre
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ancla
    if isinstance(parrafos, str):
        parrafos = [parrafos]
    for i, par in enumerate(parrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinear
        if interlineado:
            p.line_spacing = interlineado
        if espacio_despues:
            p.space_after = Pt(espacio_despues)
        runs = [par] if isinstance(par, (str, dict)) else par
        for r in runs:
            if isinstance(r, str):
                r = {"t": r}
            run = p.add_run()
            run.text = r["t"]
            f = run.font
            f.name = r.get("fuente", fuente)
            f.size = Pt(r.get("tam", tam))
            f.bold = r.get("negrita", negrita)
            f.color.rgb = rgb(r.get("color", color))
            esp = r.get("esp", espaciado)
            if esp:
                run._r.get_or_add_rPr().set("spc", str(int(esp * 100)))
    return tb


def icono(slide, nombre, x, y, lado):
    return slide.shapes.add_picture(str(ICONOS / f"{nombre}.png"), Inches(x), Inches(y), Inches(lado), Inches(lado))


def flecha(slide, x1, y1, x2, y2, color=GRIS, grosor=1.75):
    con = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    con.line.color.rgb = rgb(color)
    con.line.width = Pt(grosor)
    ln = con.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return con


def franja_superior(slide):
    # Franja naranja fina de la plantilla (todas sus láminas la usan).
    caja(slide, 0, 0, ANCHO, 0.055, relleno=NARANJA)


def pie(slide, modulo, numero, oscuro=False):
    color = GRIS if not oscuro else "8A8B8E"
    texto(slide, MARGEN, 7.02, 9, 0.25, f"IMPLEMENTACIÓN DE AWP  ·  {modulo}", tam=9, color=color, espaciado=1)
    texto(slide, ANCHO - MARGEN - 2, 7.02, 2, 0.25, f"{numero:03d}", tam=9, color=color, fuente=MONO,
          alinear=PP_ALIGN.RIGHT)


def encabezado(slide, etiqueta, titulo, oscuro=False):
    texto(slide, MARGEN, 0.42, 12, 0.3, etiqueta.upper(), tam=11, color=NARANJA, negrita=True, espaciado=3)
    texto(slide, MARGEN, 0.74, ANCHO - 2 * MARGEN, 1.0, titulo, tam=28, negrita=True,
          color=BLANCO if oscuro else NEGRO, ancla=MSO_ANCHOR.TOP, interlineado=0.95, nombre="Título")


def notas(slide, explicacion, fuentes):
    tf = slide.notes_slide.notes_text_frame
    tf.text = explicacion.strip() + "\n\nFUENTE: " + fuentes


def nueva(prs, color_fondo):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    fondo(s, color_fondo)
    franja_superior(s)
    return s


# ---------------------------------------------------------------- láminas

def portada(prs):
    s = nueva(prs, NEGRO)
    texto(s, MARGEN, 1.25, 7, 0.3, "GUÍA PRÁCTICA DE IMPLEMENTACIÓN", tam=12, color=NARANJA, negrita=True, espaciado=3)
    texto(s, MARGEN, 1.75, 7.2, 1.0, "Implementar AWP", tam=54, color=BLANCO, negrita=True, nombre="Título")
    texto(s, MARGEN, 2.85, 6.9, 1.1, ["Construir el proyecto en el orden", "en que se va a construir"],
          tam=24, color=NARANJA_CLARO, interlineado=1.05)
    texto(s, MARGEN, 4.35, 6.6, 0.8,
          [[{"t": "Advanced Work Packaging (AWP)", "negrita": True, "color": BLANCO},
            {"t": " fase por fase: roles, entregables, errores comunes y lecciones aprendidas.", "color": BORDE}]],
          tam=15, interlineado=1.15)
    texto(s, MARGEN, 6.35, 7, 0.3, "Basado en guías del CII, la COAA e Insight-AWP", tam=11, color=GRIS)

    # Jerarquía de paquetes en el estilo de "muestras de color" de la plantilla.
    x0, w = 8.35, 4.38
    niveles = [
        ("CWA", "Construction Work Area", "Área del terreno", NARANJA_OSC, "map_w"),
        ("CWP", "Construction Work Package", "Una disciplina del área", NARANJA, "helmet_w"),
        ("IWP", "Installation Work Package", "Una cuadrilla, una semana", NARANJA_CLARO, "clip_w"),
    ]
    y = 1.25
    for sigla, nombre, desc, color, ic in niveles:
        caja(s, x0, y, w, 1.05, relleno=color)
        icono(s, ic, x0 + 0.3, y + 0.3, 0.45)
        texto(s, x0 + 1.0, y + 0.18, 3.2, 0.4, sigla, tam=22, color=BLANCO, negrita=True, fuente=MONO)
        texto(s, x0 + 1.0, y + 0.6, 3.2, 0.3, nombre, tam=11, color=BLANCO)
        caja(s, x0, y + 1.05, w, 0.42, relleno=TARJETA)
        texto(s, x0 + 0.3, y + 1.05, w - 0.5, 0.42, desc, tam=11, color=BORDE, ancla=MSO_ANCHOR.MIDDLE)
        y += 1.72
    notas(s, """
Portada. Presente el propósito: esta guía explica cómo implementar Advanced Work Packaging (AWP) en un proyecto real, fase por fase.
Idea central: AWP es un proceso dirigido por Construcción. Se decide primero cómo y en qué orden se construirá la planta (Path of Construction) y, a partir de ahí, Ingeniería y Compras entregan sus productos en ese mismo orden.
Los bloques de la derecha anticipan la jerarquía que se verá en detalle: Construction Work Area (CWA) → Construction Work Package (CWP) → Installation Work Package (IWP).
Resultados que respaldan el método: +25 % de productividad y −10 % de costo total instalado (TIC), según los estudios RT-272 y RT-319 del Construction Industry Institute (CII).
""", "cii-awp-cba-education-framework-pdf.md (definición y beneficios); advanced-work-packaging-awp-quick-start-guide-pdf.md (Introducción).")


def separador(prs):
    s = nueva(prs, NEGRO)
    texto(s, MARGEN, 1.2, 4, 2.4, "06", tam=150, color=NARANJA, negrita=True, fuente=FUENTE)
    texto(s, 4.55, 1.62, 8, 0.3, "MÓDULO 6  ·  FASE 3: CONSTRUCCIÓN", tam=12, color=NARANJA, negrita=True, espaciado=3)
    texto(s, 4.55, 2.05, 8.1, 1.6, "En campo, cada capataz recibe un paquete listo para ejecutar",
          tam=36, color=BLANCO, negrita=True, interlineado=0.95, nombre="Título")
    texto(s, 4.55, 3.75, 7.8, 0.7,
          "Workface Planning (WFP), gestión de restricciones, backlog y seguimiento diario del avance.",
          tam=15, color=BORDE, interlineado=1.1)

    # Indicador de avance por fases.
    fases = [("FASE 1", "Planificación · FEL 2"), ("FASE 2", "Ingeniería y compras · FEL 3"),
             ("FASE 3", "Construcción"), ("FASE 4", "Puesta en marcha")]
    x, w, gap, y = MARGEN, 2.85, 0.24, 5.35
    for i, (f, n) in enumerate(fases):
        activo = i == 2
        caja(s, x, y, w, 0.95, relleno=NARANJA if activo else TARJETA,
             borde=None if activo else "333333")
        texto(s, x + 0.25, y + 0.17, w - 0.4, 0.25, f, tam=10, negrita=True, espaciado=2,
              color=BLANCO if activo else NARANJA)
        texto(s, x + 0.25, y + 0.47, w - 0.4, 0.3, n, tam=13, negrita=activo,
              color=BLANCO if activo else BORDE)
        x += w + gap
    pie(s, "M6 · CONSTRUCCIÓN Y WORKFACE PLANNING", 59, oscuro=True)
    notas(s, """
Separador del Módulo 6. Hasta aquí, Construcción definió el orden (Fase 1) e Ingeniería y Compras entregaron sus paquetes en ese orden (Fase 2). Ahora el plan llega al frente de trabajo.
En este módulo se verá el proceso de Workface Planning (WFP): cómo se dividen los Construction Work Packages (CWP) en Installation Work Packages (IWP), cómo se eliminan las restricciones antes de liberar un paquete, cómo se mantiene un backlog y cómo se registra el avance.
La barra inferior indica en qué fase del proyecto estamos; se repite en todos los separadores de fase.
""", "advanced-work-packaging-awp-quick-start-guide-pdf.md (secciones 5-7); 3-workface-planning-procedure-insight-awp-2017-pdf.md.")


def jerarquia(prs):
    s = nueva(prs, FONDO)
    encabezado(s, "Módulo 2 · Qué es AWP", "Todo se ordena en una jerarquía: del área al paquete que recibe el capataz")

    def nodo(x, y, w, h, sigla, nombre, dato, relleno, ic, txt=BLANCO, sub=None):
        caja(s, x, y, w, h, relleno=relleno)
        icono(s, ic, x + 0.22, y + (h - 0.42) / 2, 0.42)
        texto(s, x + 0.85, y + 0.14, w - 1.0, 0.35, sigla, tam=18, negrita=True, color=txt, fuente=MONO)
        texto(s, x + 0.85, y + 0.48, w - 1.0, 0.25, nombre, tam=10.5, color=txt)
        texto(s, x + 0.85, y + 0.74, w - 1.0, 0.25, dato, tam=10.5, color=sub or txt, negrita=True)

    cx = ANCHO / 2
    w, h = 3.5, 1.08
    # Nivel 1: CWA
    nodo(cx - w / 2, 1.95, w, h, "CWA", "Construction Work Area", "Área del terreno · nivel 2", NEGRO, "map_w",
         sub=NARANJA_CLARO)
    # Nivel 2: EWP → CWP ← PWP
    y2 = 3.55
    nodo(cx - w / 2, y2, w, h, "CWP", "Construction Work Package", "1 disciplina · < 40 000 HH · nivel 3",
         NARANJA, "helmet_w")
    xe = MARGEN + 0.15
    nodo(xe, y2, w, h, "EWP", "Engineering Work Package", "Planos, specs, lista de materiales", AZUL, "ruler_w",
         sub=AZUL_PALIDO)
    xp = ANCHO - MARGEN - 0.15 - w
    nodo(xp, y2, w, h, "PWP", "Procurement Work Package", "Materiales y equipos del CWP", AZUL_MEDIO, "truck_w")
    flecha(s, xe + w + 0.08, y2 + h / 2, cx - w / 2 - 0.1, y2 + h / 2, color=AZUL)
    flecha(s, xp - 0.08, y2 + h / 2, cx + w / 2 + 0.1, y2 + h / 2, color=AZUL_MEDIO)
    flecha(s, cx, 1.95 + h + 0.06, cx, y2 - 0.08, color=GRIS)

    # Nivel 3: IWPs
    y3 = 5.15
    wi, gi = 2.5, 0.25
    total = 3 * wi + 2 * gi
    x = cx - total / 2
    for i in range(3):
        caja(s, x, y3, wi, 0.95, relleno=BLANCO, borde=BORDE)
        icono(s, "clip_o", x + 0.2, y3 + 0.26, 0.42)
        texto(s, x + 0.8, y3 + 0.14, wi - 0.9, 0.35, f"IWP-0{i + 1}", tam=15, negrita=True, color=NEGRO, fuente=MONO)
        texto(s, x + 0.8, y3 + 0.52, wi - 0.9, 0.3, "1 cuadrilla · 1 semana", tam=10.5, color=GRIS)
        flecha(s, cx, y2 + h + 0.06, x + wi / 2, y3 - 0.08, color=GRIS, grosor=1.25)
        x += wi + gi
    texto(s, cx - total / 2, y3 + 1.12, total, 0.3,
          [[{"t": "Installation Work Package (IWP): ", "negrita": True, "color": NEGRO},
            {"t": "lo que recibe el capataz · nivel 5"}]], tam=11, color=GRIS, alinear=PP_ALIGN.CENTER)
    texto(s, cx + total / 2 + 0.3, y3 + 0.05, 2.7, 0.9,
          [[{"t": "Regla base", "negrita": True, "color": NARANJA}], "1 EWP = 1 PWP = 1 CWP"],
          tam=11, color=TEXTO, interlineado=1.1)
    pie(s, "M2 · QUÉ ES AWP", 16)
    notas(s, """
La jerarquía de paquetes es el esqueleto de AWP. Léala de arriba hacia abajo:
• Construction Work Area (CWA): división lógica del plano de implantación. Es una actividad de nivel 2 del cronograma.
• Construction Work Package (CWP): una sola disciplina dentro de una CWA, con menos de 40 000 horas-hombre. Es una actividad de nivel 3.
• Engineering Work Package (EWP): contiene toda la ingeniería que necesita un CWP (alcance, planos, datos de proveedor, lista de materiales y especificaciones).
• Procurement Work Package (PWP): agrupa los materiales y equipos de un CWP. El CII todavía lo trata como paquete en definición (RT-363).
• Installation Work Package (IWP): trabajo de un capataz y su cuadrilla en unos 7 días, formado por planos completos. Es una actividad de nivel 5.
Regla base: 1 EWP = 1 PWP = 1 CWP. Omega 365 y el curso 2023 admiten excepciones, como materiales a granel o varios EWP para un mismo CWP.
Advierta que el Glosario del CII ubica el IWP en el nivel 4. Cada proyecto debe fijar su criterio en el procedimiento.
""", "1-advanced-work-packaging-procedure-insight-awp-2017-pdf.md (§2 Definiciones); advanced-work-packaging-a-recommended-practice-for-capital-projects-omega-365-blogs.md (jerarquía de paquetes); awp-acronyms-definitions-pdf.md; INDICE.md §3.")


def proceso(prs):
    s = nueva(prs, FONDO)
    encabezado(s, "Módulo 6 · Gestión de restricciones",
               "Las restricciones se liberan en cuatro pasos, con semanas de anticipación")
    pasos = [
        ("clip_o", "IWP iniciado", "El planificador crea el paquete a partir del CWP", 6, 12),
        ("search_o", "Restricciones identificadas", "Documentos, materiales, andamios, equipos, permisos", 4, 10),
        ("usercheck_o", "Restricciones asignadas", "Cada una con dueño, fecha y prioridad", 3, 8),
        ("unlock_o", "Restricciones liberadas", "El IWP pasa al backlog, listo para campo", 2, 4),
    ]
    x, w, gap, y, h = MARGEN, 2.78, 0.3, 2.05, 3.55
    for i, (ic, tit, desc, peq, mega) in enumerate(pasos):
        caja(s, x, y, w, h, relleno=BLANCO, borde=BORDE)
        num = caja(s, x + 0.3, y + 0.3, 0.55, 0.55, relleno=NARANJA, forma=MSO_SHAPE.OVAL)
        texto(s, x + 0.3, y + 0.3, 0.55, 0.55, str(i + 1), tam=18, negrita=True, color=BLANCO,
              alinear=PP_ALIGN.CENTER, ancla=MSO_ANCHOR.MIDDLE)
        icono(s, ic, x + w - 0.8, y + 0.33, 0.48)
        texto(s, x + 0.3, y + 1.08, w - 0.55, 0.65, tit, tam=15, negrita=True, color=NEGRO, interlineado=0.95)
        texto(s, x + 0.3, y + 1.75, w - 0.55, 0.7, desc, tam=11.5, color=GRIS, interlineado=1.05)
        caja(s, x + 0.3, y + 2.55, w - 0.6, 0.012, relleno=BORDE)
        for j, (etq, val) in enumerate((("Pequeño / parada", peq), ("Megaproyecto", mega))):
            yy = y + 2.68 + j * 0.36
            texto(s, x + 0.3, yy, 1.6, 0.3, etq, tam=10, color=GRIS, ancla=MSO_ANCHOR.MIDDLE)
            texto(s, x + w - 1.3, yy, 1.0, 0.3, f"−{val} sem", tam=12, negrita=True,
                  color=NARANJA if j == 0 else AZUL, fuente=MONO, alinear=PP_ALIGN.RIGHT, ancla=MSO_ANCHOR.MIDDLE)
        if i < 3:
            flecha(s, x + w + 0.04, y + h / 2, x + w + gap - 0.04, y + h / 2, color=NARANJA, grosor=2)
        x += w + gap
    caja(s, MARGEN, 5.9, ANCHO - 2 * MARGEN, 0.75, relleno=NEGRO)
    icono(s, "flag_o", MARGEN + 0.3, 6.07, 0.4)
    texto(s, MARGEN + 0.95, 5.9, 10.8, 0.75,
          [[{"t": "Regla de oro: ", "negrita": True, "color": NARANJA},
            {"t": "ningún IWP va a campo mientras tenga una restricción abierta.", "color": BLANCO}]],
          tam=15, ancla=MSO_ANCHOR.MIDDLE)
    pie(s, "M6 · CONSTRUCCIÓN Y WORKFACE PLANNING", 66)
    notas(s, """
El Education Framework del CII propone un calendario típico de restricciones para cada IWP, medido en semanas antes de su ejecución:
• Proyecto pequeño, parada de planta o turnaround: IWP iniciado 6 semanas antes; restricciones identificadas a las 4; asignadas a las 3; liberadas a las 2 semanas, cuando el paquete se libera.
• Proyecto grande, mega o giga: 12, 10, 8 y 4 semanas, respectivamente.
Críticas: documentos, materiales y andamios. Secundarias: equipos de construcción, controles, seguridad, calidad y personal, que el equipo resuelve internamente en poco tiempo (Quick Start Guide).
Buenas prácticas: una lista única de restricciones abiertas; dueño, fecha y prioridad para cada una; flujo de aprobación para cerrarlas; reunión estándar de revisión; árbol de decisión para paquetes con restricciones abiertas; reporte de estado.
La regla de oro de la Quick Start Guide dice que los IWP no se liberan a campo hasta que estén libres de restricciones y listos para ejecutarse. El Workface Planner recoge las firmas en la página de restricciones del IWP.
Nota: el curso 2023 indica liberar 2-3 semanas antes, y la Quick Start pide enviar la lista de materiales 8 semanas antes. Ajuste los plazos al proyecto.
""", "cii-awp-cba-education-framework-pdf.md (Typical IWP Constraint Schedule by Project Size / Type; Constraint Management Best Practices); advanced-work-packaging-awp-quick-start-guide-pdf.md (§7 Constraint Removal).")


def comparacion(prs):
    s = nueva(prs, FONDO)
    encabezado(s, "Módulo 5 · Fase 2: Ingeniería",
               "Tradicional: se diseña lo disponible. AWP: se diseña lo que se construirá primero")
    paneles = [
        ("ENFOQUE TRADICIONAL", GRIS_MEDIO, "xmark_g",
         [("Secuencia de ingeniería", "Ingeniería avanza según su propia lógica"),
          ("Liberación a construcción", "Se emiten planos a medida que se terminan"),
          ("Se ejecuta el alcance disponible", "El campo trabaja donde hay información")],
         "Frentes incompletos, esperas y retrabajo"),
        ("ENFOQUE AWP", NARANJA, "check_o",
         [("Secuencia de construcción", "El Path of Construction (POC) fija el orden"),
          ("Secuencia de ingeniería", "Los EWP se entregan en el orden del POC"),
          ("Se ejecuta según el plan acordado", "Cada CWP llega completo a campo")],
         "Paquetes completos, en orden y a tiempo"),
    ]
    w, gap, y, h = 5.9, 0.33, 2.0, 4.7
    x = MARGEN
    for tit, color, ic, pasos, res in paneles:
        caja(s, x, y, w, h, relleno=BLANCO, borde=BORDE)
        caja(s, x, y, w, 0.55, relleno=color)
        texto(s, x + 0.3, y, w - 0.6, 0.55, tit, tam=12, negrita=True, color=BLANCO, espaciado=2,
              ancla=MSO_ANCHOR.MIDDLE)
        yy = y + 0.8
        for k, (p, d) in enumerate(pasos):
            caja(s, x + 0.3, yy, 0.42, 0.42, relleno=color, forma=MSO_SHAPE.OVAL)
            texto(s, x + 0.3, yy, 0.42, 0.42, str(k + 1), tam=13, negrita=True, color=BLANCO,
                  alinear=PP_ALIGN.CENTER, ancla=MSO_ANCHOR.MIDDLE)
            texto(s, x + 0.95, yy - 0.02, w - 1.25, 0.3, p, tam=14, negrita=True, color=NEGRO)
            texto(s, x + 0.95, yy + 0.3, w - 1.25, 0.3, d, tam=11.5, color=GRIS)
            if k < 2:
                flecha(s, x + 0.51, yy + 0.47, x + 0.51, yy + 0.83, color=color, grosor=1.5)
            yy += 0.9
        caja(s, x + 0.3, y + h - 0.95, w - 0.6, 0.7, relleno=FONDO)
        icono(s, ic, x + 0.5, y + h - 0.8, 0.4)
        texto(s, x + 1.1, y + h - 0.95, w - 1.5, 0.7, res, tam=13.5, negrita=True, color=NEGRO,
              ancla=MSO_ANCHOR.MIDDLE)
        x += w + gap
    pie(s, "M5 · INGENIERÍA DE DETALLE Y COMPRAS", 49)
    notas(s, """
La diferencia no está en cuánto trabajo hace Ingeniería, sino en el orden en que lo entrega. El curso AWP 2023 lo resume: la intención no es crear trabajo extra para Ingeniería ni cambiar cómo diseña, sino alinear su secuencia con la de Construcción.
Enfoque tradicional: Ingeniería define su propia secuencia, libera documentos a Construcción a medida que los termina, y Construcción ejecuta el alcance que tenga disponible. El resultado son frentes de trabajo incompletos y cuadrillas esperando información.
Enfoque AWP: Construcción define primero el Path of Construction (POC). Ingeniería usa el POC como base de su Path of Engineering y entrega Engineering Work Packages (EWP) completos por Construction Work Area (CWA). Construcción ejecuta según el plan acordado.
Beneficio adicional: gestionar la ingeniería por CWA crea una serie de miniproyectos con entregables más pequeños, lo que permite detectar atrasos antes.
""", "awp-2023-integrado-pdf.md (Lección 3: AWP and Engineering, 'Understanding the Path of Engineering', 'Managing Engineering by CWA').")


def tabla(prs):
    s = nueva(prs, FONDO)
    encabezado(s, "Módulo 4 · Fase 1: Planificación", "El cronograma se escalona en 7 niveles, del año a la hora")
    filas = [
        ("1", "Planta", "Año", ""),
        ("2", "Construction Work Area", "6 meses", "CWA"),
        ("3", "Construction Work Package", "3 meses", "CWP"),
        ("4", "Código de función", "Mes", ""),
        ("5", "Installation Work Package", "Semana", "IWP"),
        ("6", "Plan diario del capataz", "Día", ""),
        ("7", "Tarea", "Hora", ""),
    ]
    cols = [1.2, 5.2, 2.4, 2.6]
    x0, y0, hh, hf = MARGEN + 0.3, 1.95, 0.5, 0.52
    enc = ["NIVEL", "UNIDAD DE PLANIFICACIÓN", "HORIZONTE", "PAQUETE AWP"]
    shp = s.shapes.add_table(len(filas) + 1, 4, Inches(x0), Inches(y0), Inches(sum(cols)),
                             Inches(hh + hf * len(filas)))
    tbl = shp.table
    # Estilo propio: sin bandas del estilo por defecto.
    tblPr = tbl._tbl.tblPr
    tblPr.set("bandRow", "0")
    tblPr.set("firstRow", "0")
    for i, wcol in enumerate(cols):
        tbl.columns[i].width = Inches(wcol)
    tbl.rows[0].height = Inches(hh)
    for r in range(1, len(filas) + 1):
        tbl.rows[r].height = Inches(hf)

    def celda(c, txt, tam, color, negrita, relleno, fuente=FUENTE, alinear=PP_ALIGN.LEFT, esp=0):
        c.fill.solid()
        c.fill.fore_color.rgb = rgb(relleno)
        c.margin_left = Inches(0.2)
        c.margin_right = Inches(0.1)
        c.margin_top = c.margin_bottom = 0
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = c.text_frame
        tf.text = ""
        p = tf.paragraphs[0]
        p.alignment = alinear
        run = p.add_run()
        run.text = txt
        run.font.size = Pt(tam)
        run.font.bold = negrita
        run.font.name = fuente
        run.font.color.rgb = rgb(color)
        if esp:
            run._r.get_or_add_rPr().set("spc", str(esp * 100))

    for j, t in enumerate(enc):
        celda(tbl.cell(0, j), t, 10.5, BLANCO, True, NEGRO, esp=1)
    for i, (niv, uni, hor, paq) in enumerate(filas, start=1):
        destacado = paq in ("CWP", "IWP")
        base = "FFE3CF" if destacado else (BLANCO if i % 2 else "F7F7F7")
        celda(tbl.cell(i, 0), niv, 16, NARANJA if destacado else GRIS, True, base, fuente=MONO)
        celda(tbl.cell(i, 1), uni, 14, NEGRO, destacado, base)
        celda(tbl.cell(i, 2), hor, 14, TEXTO, False, base)
        celda(tbl.cell(i, 3), paq, 14, NARANJA_OSC if destacado else AZUL, True, base, fuente=MONO)

    # Nota al pie de la tabla
    yb = y0 + hh + hf * len(filas) + 0.25
    icono(s, "layers_o", x0, yb + 0.02, 0.32)
    texto(s, x0 + 0.5, yb, 10.5, 0.4,
          [[{"t": "El CWP se divide en IWP: ", "negrita": True, "color": NEGRO},
            {"t": "el planificador entrega esa secuencia al programador como cronograma de nivel 5."}]],
          tam=12, color=TEXTO, ancla=MSO_ANCHOR.MIDDLE)
    pie(s, "M4 · FASE 1: PLANIFICACIÓN PRELIMINAR", 45)
    notas(s, """
Los procedimientos de Insight-AWP (2017) definen siete niveles de cronograma. Cada nivel tiene su unidad de planificación y su horizonte:
Nivel 1: planta, año. Nivel 2: Construction Work Area (CWA), 6 meses. Nivel 3: Construction Work Package (CWP), 3 meses. Nivel 4: código de función, mes. Nivel 5: Installation Work Package (IWP), semana. Nivel 6: plan diario, día. Nivel 7: tarea, hora.
La clave práctica es que el cronograma del proyecto se lleva hasta el nivel 3 con CWP. Luego el Workface Planner lo baja al nivel 5 dividiendo cada CWP en IWP y entregando la secuencia al programador. Según la Quick Start Guide, el cronograma de nivel 5 solo debe contener CWP, IWP e hitos.
Advertencia: el Glosario del CII ubica el IWP en el nivel 4. Defina el criterio en el procedimiento del proyecto y aplíquelo de forma consistente.
""", "1-advanced-work-packaging-procedure-insight-awp-2017-pdf.md (§2 Definiciones, niveles de cronograma); advanced-work-packaging-awp-quick-start-guide-pdf.md (§6 The Project Schedule); INDICE.md §3.")


def dato_clave(prs):
    s = nueva(prs, FONDO)
    encabezado(s, "Módulo 10 · Medición", "El tool time pasa de 37 % a 46 % con AWP")
    # Cifras grandes
    texto(s, MARGEN, 2.0, 3.2, 1.3, "37 %", tam=80, negrita=True, color=GRIS)
    texto(s, MARGEN, 3.35, 3.4, 0.6, [[{"t": "Proyectos tradicionales", "negrita": True, "color": NEGRO}],
                                     "3,7 h productivas por jornada de 10 h"], tam=12, color=GRIS)
    flecha(s, 3.95, 2.72, 4.95, 2.72, color=NARANJA, grosor=3)
    texto(s, 5.2, 2.0, 3.4, 1.3, "46 %", tam=80, negrita=True, color=NARANJA)
    texto(s, 5.2, 3.35, 3.4, 0.6, [[{"t": "Proyectos con AWP", "negrita": True, "color": NEGRO}],
                                  "4,6 h productivas por jornada de 10 h"], tam=12, color=GRIS)

    # Barras de 10 h
    yb = 4.55
    texto(s, MARGEN, yb - 0.4, 6, 0.3, "JORNADA DE 10 HORAS", tam=10, negrita=True, color=GRIS, espaciado=2)
    wbar = 8.0
    for k, (etq, pct, col) in enumerate((("Tradicional", 0.37, GRIS), ("AWP", 0.46, NARANJA))):
        y = yb + k * 0.72
        texto(s, MARGEN, y, 1.3, 0.5, etq, tam=12, negrita=True, color=NEGRO, ancla=MSO_ANCHOR.MIDDLE)
        caja(s, MARGEN + 1.4, y, wbar * pct, 0.5, relleno=col)
        caja(s, MARGEN + 1.4 + wbar * pct, y, wbar * (1 - pct), 0.5, relleno="E0E0E0")
        texto(s, MARGEN + 1.55, y, 2, 0.5, f"{pct * 10:.1f} h".replace(".", ","), tam=12, negrita=True,
              color=BLANCO, ancla=MSO_ANCHOR.MIDDLE, fuente=MONO)
        texto(s, MARGEN + 1.4 + wbar * pct + 0.15, y, 3, 0.5,
              f"{(1 - pct) * 10:.1f} h sin herramientas".replace(".", ","), tam=11, color=GRIS,
              ancla=MSO_ANCHOR.MIDDLE)
    # Tarjeta lateral
    xc = 10.35
    caja(s, xc, 2.0, ANCHO - MARGEN - xc, 3.77, relleno=NEGRO)
    icono(s, "wrench_o", xc + 0.35, 2.35, 0.5)
    texto(s, xc + 0.35, 3.1, 1.95, 0.9, "+25 %", tam=40, negrita=True, color=NARANJA)
    texto(s, xc + 0.35, 3.95, 1.95, 1.2, "de productividad en campo (CII)", tam=13, color=BLANCO, interlineado=1.1)
    texto(s, xc + 0.35, 4.85, 1.95, 0.8, "Mídalo con estudios de tool time cada 2-3 meses",
          tam=11, color=BORDE, interlineado=1.1)
    pie(s, "M10 · MEDICIÓN: KPIs, CONTROLES Y PRODUCTIVIDAD", 101)
    notas(s, """
El tool time, o tiempo en herramientas, es el porcentaje de la jornada en que el trabajador está instalando. Es el indicador principal de productividad en campo.
Según el Education Framework del CII, en proyectos tradicionales es del 37 % (3,7 h de una jornada de 10 h) y con AWP sube al 46 % (4,6 h). El curso 2023 cita 47 %. El resto del tiempo se pierde esperando información, materiales, herramientas o accesos, desplazándose o en pausas.
0,9 h adicionales de trabajo efectivo por persona y día equivalen a un aumento de productividad de alrededor del 24-25 %, coherente con el +25 % de productividad y el −10 % de costo total instalado (TIC) que reportan los estudios RT-272 y RT-319 del CII.
Cómo medirlo: el Procedimiento 3.0 recomienda estudios de tool time hechos por terceros cada 2 o 3 meses. Estos estudios son una "nota" para la gerencia, no para los trabajadores: miden si se entregan información, herramientas, materiales y acceso a tiempo.
""", "cii-awp-cba-education-framework-pdf.md (Benefits & Value of AWP: 37 % vs 46 %); awp-2023-integrado-pdf.md (Lección 2: Tool Time 37 % / 47 %); 3-workface-planning-procedure-insight-awp-2017-pdf.md (§18 Tool Time Studies).")


def cierre(prs):
    s = nueva(prs, NEGRO)
    texto(s, MARGEN, 0.95, 10, 0.3, "EN RESUMEN", tam=12, negrita=True, color=NARANJA, espaciado=3)
    texto(s, MARGEN, 1.35, 11.5, 1.4, "Construya en el orden correcto y la productividad llega sola",
          tam=40, negrita=True, color=BLANCO, interlineado=0.95, nombre="Título")
    ideas = [
        ("helmet_o", "Construcción define el orden", "Path of Construction (POC) y CWP"),
        ("ruler_o", "Ingeniería y Compras lo siguen", "EWP y PWP en la secuencia del POC"),
        ("clip_o", "El campo ejecuta sin restricciones", "IWP listos, backlog y avance diario"),
    ]
    w, gap, y = 3.83, 0.32, 3.2
    x = MARGEN
    for ic, t, d in ideas:
        caja(s, x, y, w, 2.15, relleno=TARJETA, borde="333333")
        icono(s, ic, x + 0.35, y + 0.3, 0.5)
        texto(s, x + 0.35, y + 0.9, w - 0.6, 0.62, t, tam=15, negrita=True, color=BLANCO,
              ancla=MSO_ANCHOR.BOTTOM, interlineado=0.95)
        texto(s, x + 0.35, y + 1.6, w - 0.6, 0.35, d, tam=12, color=BORDE)
        x += w + gap
    texto(s, MARGEN, 5.8, 6, 0.6, "¿Preguntas?", tam=28, negrita=True, color=NARANJA_CLARO)
    texto(s, 6.9, 5.85, ANCHO - MARGEN - 6.9, 0.6, "Siguiente paso: nombrar al AWP Champion del proyecto",
          tam=14, color=BORDE, alinear=PP_ALIGN.RIGHT, ancla=MSO_ANCHOR.MIDDLE)
    pie(s, "CIERRE", 122, oscuro=True)
    notas(s, """
Cierre. Resuma el hilo conductor de toda la presentación en tres ideas:
1. Construcción define el orden. El Path of Construction (POC) se construye en sesiones de planificación interactiva, se define en FEL 2 y se congela en FEL 3. De él salen las Construction Work Areas (CWA) y los Construction Work Packages (CWP).
2. Ingeniería y Compras siguen ese orden. Los Engineering Work Packages (EWP) y Procurement Work Packages (PWP) se entregan en la secuencia del POC (1 EWP = 1 PWP = 1 CWP).
3. El campo ejecuta sin restricciones. Cada capataz recibe un Installation Work Package (IWP) libre de restricciones, tomado de un backlog, y registra el avance a diario.
Siguiente paso concreto: nombrar un AWP Champion dedicado a tiempo completo. Es el primer componente de la Quick Start Guide y la lección 8 del curso 2023.
Abra la ronda de preguntas.
""", "advanced-work-packaging-awp-quick-start-guide-pdf.md (§1 y §8); awp-2023-integrado-pdf.md (AWP Related Lessons Learned, lección 8); síntesis de la presentación.")


def main():
    prs = Presentation(str(PLANTILLA))
    # Quitar las 20 láminas de catálogo de la plantilla; se conservan patrón, diseño y tema.
    lista = prs.slides._sldIdLst
    for sld in list(lista):
        prs.part.drop_rel(sld.rId)
        lista.remove(sld)
    # El patrón de notas de la plantilla no tiene marcador de cuerpo: se reemplaza por el estándar.
    pres = prs.part._element
    nm = pres.find(qn("p:notesMasterIdLst"))
    if nm is not None:
        prs.part.drop_rel(nm[0].get(qn("r:id")))
        pres.remove(nm)
    for fn in (portada, separador, jerarquia, proceso, comparacion, tabla, dato_clave, cierre):
        fn(prs)
    prs.core_properties.title = "Implementar AWP: muestra de diseño"
    prs.core_properties.language = "es-PE"
    prs.save(str(SALIDA))
    print("Guardado:", SALIDA)


if __name__ == "__main__":
    main()
