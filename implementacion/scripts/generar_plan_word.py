"""Genera implementacion/Plan_Implementacion_AWP.docx a partir de las páginas
de la sección "Implementación" de la web (docs/implementacion/*.md).

Así la web y el Word tienen siempre el mismo contenido. Los diagramas Mermaid
se incluyen como imágenes PNG guardadas en implementacion/scripts/diagramas/.
Si falta alguna imagen (o se usa --renderizar), se genera con mermaid-cli
(`mmdc`), que debe estar instalado; la ruta se puede indicar con la variable
de entorno MMDC y la configuración de Puppeteer con MMDC_PUPPETEER.

Uso:
    python implementacion/scripts/generar_plan_word.py [--renderizar]
"""

import hashlib
import os
import re
import struct
import subprocess
import sys
from pathlib import Path

import yaml
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
import estilo_word as ew  # noqa: E402

RAIZ = Path(__file__).resolve().parents[2]
DOCS = RAIZ / "docs"
DIAGRAMAS = Path(__file__).resolve().parent / "diagramas"
SALIDA = RAIZ / "implementacion" / "Plan_Implementacion_AWP.docx"

# Título del capítulo en el Word cuando difiere del título de la página web.
TITULOS = {"implementacion/index.md": "Objetivo, alcance y beneficios esperados"}


# --------------------------------------------------------------------------
# Lectura del menú de MkDocs
# --------------------------------------------------------------------------

class _Cargador(yaml.SafeLoader):
    pass


# mkdocs.yml usa etiquetas !!python/name: que no hace falta resolver aquí.
_Cargador.add_multi_constructor("tag:yaml.org,2002:python/", lambda l, s, n: None)


def paginas_implementacion():
    config = yaml.load((RAIZ / "mkdocs.yml").read_text(encoding="utf-8"), Loader=_Cargador)
    for seccion in config["nav"]:
        if isinstance(seccion, dict) and "Implementación" in seccion:
            return [list(e.items())[0] for e in seccion["Implementación"]]
    raise SystemExit("No se encontró la sección 'Implementación' en mkdocs.yml")


# --------------------------------------------------------------------------
# Diagramas
# --------------------------------------------------------------------------

def imagen_diagrama(codigo, renderizar):
    DIAGRAMAS.mkdir(exist_ok=True)
    clave = hashlib.sha1(codigo.encode("utf-8")).hexdigest()[:12]
    png = DIAGRAMAS / f"diagrama_{clave}.png"
    if png.exists() and not renderizar:
        return png
    mmd = DIAGRAMAS / f"diagrama_{clave}.mmd"
    mmd.write_text(codigo, encoding="utf-8")
    cmd = [os.environ.get("MMDC", "mmdc"), "-i", str(mmd), "-o", str(png), "-s", "2", "-b", "white"]
    if os.environ.get("MMDC_PUPPETEER"):
        cmd += ["-p", os.environ["MMDC_PUPPETEER"]]
    try:
        subprocess.run(cmd, check=True, capture_output=True, timeout=120)
    except (OSError, subprocess.SubprocessError) as error:
        print(f"Aviso: no se pudo renderizar un diagrama ({error}).")
        return None
    finally:
        mmd.unlink(missing_ok=True)
    return png


# --------------------------------------------------------------------------
# Texto en línea: negrita, cursiva, código y enlaces
# --------------------------------------------------------------------------

PATRON = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))")


def texto_en_linea(parrafo, texto, tam=None, negrita=False):
    texto = texto.replace("<br/>", " ").replace("<br>", " ")
    # Íconos de la web (:material-...:) que no tienen equivalente en Word.
    texto = re.sub(r":(material|octicons)-[a-z0-9-]+:\s*", "", texto)
    for trozo in PATRON.split(texto):
        if not trozo:
            continue
        estilo = {}
        if trozo.startswith("**") and trozo.endswith("**"):
            trozo, estilo = trozo[2:-2], {"bold": True}
        elif trozo.startswith("`") and trozo.endswith("`"):
            trozo, estilo = trozo[1:-1], {"code": True}
        elif trozo.startswith("[") and "](" in trozo:
            trozo = trozo[1 : trozo.index("](")]
            estilo = {"link": True}
        elif trozo.startswith("*") and trozo.endswith("*") and len(trozo) > 2:
            trozo, estilo = trozo[1:-1], {"italic": True}
        # Negritas o cursivas anidadas dentro de un enlace o cursiva.
        trozo = trozo.replace("**", "")
        run = parrafo.add_run(trozo)
        run.bold = negrita or estilo.get("bold", False)
        run.italic = estilo.get("italic", False)
        if estilo.get("code"):
            run.font.name = "Consolas"
            run.font.size = Pt(tam - 0.5 if tam else 9.5)
        elif tam:
            run.font.size = Pt(tam)
        if estilo.get("link"):
            run.font.color.rgb = ew.AZUL
            run.underline = True


# --------------------------------------------------------------------------
# Conversión de Markdown a Word
# --------------------------------------------------------------------------

def celdas(linea):
    return [c.strip() for c in linea.strip().strip("|").split("|")]


def anchos_para(encabezados, filas, total=17.0, tam=9):
    """Anchos de columna: cada columna recibe al menos el ancho de su palabra
    más larga y el resto se reparte según la longitud del texto."""
    cm_por_caracter = 0.19 if tam >= 9 else 0.17

    def limpio(x):
        return re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", x).replace("**", "").replace("`", "")

    minimos, pesos = [], []
    for i in range(len(encabezados)):
        col = [f[i] if i < len(f) else "" for f in filas]
        palabras = [p for x in col for p in limpio(x).split()] or [""]
        palabra_enc = max((len(p) for p in limpio(encabezados[i]).split()), default=4)
        minimos.append(min(4.0, max(max(len(p) for p in palabras), palabra_enc, 3) * cm_por_caracter + 0.45))
        pesos.append(min(max((len(limpio(x)) for x in col + [encabezados[i]]), default=6), 70))
    resto = max(total - sum(minimos), 0)
    suma = sum(pesos) or 1
    anchos = [m + resto * p / suma for m, p in zip(minimos, pesos)]
    factor = total / sum(anchos)
    return [a * factor for a in anchos]


def quitar_numero(titulo):
    return re.sub(r"^\d+\.\s+", "", titulo)


class Conversor:
    def __init__(self, doc, renderizar):
        self.doc = doc
        self.renderizar = renderizar
        self.cap = 0
        self.sub = 0
        self.subsub = 0
        self.figura = 0
        self.tabla_n = 0

    def titulo(self, nivel, texto):
        texto = quitar_numero(texto)
        if nivel == 1:
            self.cap += 1
            self.sub = self.subsub = 0
            numero = f"{self.cap}."
        elif nivel == 2:
            self.sub += 1
            self.subsub = 0
            numero = f"{self.cap}.{self.sub}"
        else:
            self.subsub += 1
            numero = f"{self.cap}.{self.sub}.{self.subsub}"
        self.doc.add_heading(f"{numero}  {texto}", level=min(nivel, 3))

    def figura_mermaid(self, codigo):
        png = imagen_diagrama(codigo, self.renderizar)
        if not png:
            ew.recuadro(self.doc, "Diagrama", "Ver el diagrama en la versión web del plan.")
            return
        # Tamaño del PNG leído de su cabecera (sin depender de Pillow).
        cabecera = png.read_bytes()[16:24]
        ancho, alto = struct.unpack(">II", cabecera)
        ancho_cm = 16.5
        if alto / ancho * ancho_cm > 19:
            ancho_cm = 19 * ancho / alto
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(png), width=Cm(ancho_cm))
        self.figura += 1
        titulo = re.search(r"title\s+\"?([^\"\n]+)\"?", codigo)
        leyenda = self.doc.add_paragraph()
        leyenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = leyenda.add_run(f"Figura {self.figura}" + (f". {titulo.group(1).strip()}" if titulo else ""))
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    def tabla(self, lineas):
        encabezados = celdas(lineas[0])
        filas = [celdas(l) for l in lineas[2:]]
        filas = [f + [""] * (len(encabezados) - len(f)) for f in filas]
        tam = 9 if len(encabezados) <= 5 else 8
        valores = [
            [(lambda par, v=v: texto_en_linea(par, v, tam)) for v in fila] for fila in filas
        ]
        encabezados_limpios = [re.sub(r"[*`]", "", e) for e in encabezados]
        ew.tabla_datos(self.doc, encabezados_limpios, valores, anchos_para(encabezados, filas, tam=tam), tam)

    def recuadro(self, titulo, lineas):
        partes = []
        for l in lineas:
            l = l.strip()
            if not l:
                continue
            if l.startswith("- "):
                l = "• " + l[2:]
            partes.append(lambda par, v=l: texto_en_linea(par, v, 10))
        ew.recuadro(self.doc, titulo, partes)

    def pagina(self, texto, titulo_capitulo=None):
        lineas = texto.splitlines()
        i = 0
        while i < len(lineas):
            linea = lineas[i]
            if linea.startswith("```mermaid"):
                j = i + 1
                while not lineas[j].startswith("```"):
                    j += 1
                self.figura_mermaid("\n".join(lineas[i + 1 : j]))
                i = j + 1
                continue
            m = re.match(r"^(#{1,4})\s+(.*)", linea)
            if m:
                nivel = len(m.group(1))
                texto_t = titulo_capitulo if (nivel == 1 and titulo_capitulo) else m.group(2)
                self.titulo(nivel, texto_t)
                i += 1
                continue
            if linea.startswith("!!!"):
                m = re.match(r'^!!!\s+\w+(?:\s+"([^"]*)")?', linea)
                titulo = m.group(1) or "Nota"
                j = i + 1
                cuerpo = []
                while j < len(lineas) and (lineas[j].startswith("    ") or not lineas[j].strip()):
                    cuerpo.append(lineas[j])
                    j += 1
                self.recuadro(titulo, cuerpo)
                i = j
                continue
            if linea.startswith("|"):
                j = i
                while j < len(lineas) and lineas[j].startswith("|"):
                    j += 1
                self.tabla(lineas[i:j])
                i = j
                continue
            m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", linea)
            if m:
                numerada = m.group(2)[0].isdigit()
                estilo = "List Number" if numerada else "List Bullet"
                p = self.doc.add_paragraph(style=estilo)
                if len(m.group(1)) >= 4:
                    p.paragraph_format.left_indent = Cm(1.6)
                texto_en_linea(p, m.group(3))
                i += 1
                continue
            if linea.strip():
                p = self.doc.add_paragraph()
                cursiva = linea.startswith("*") and linea.rstrip().endswith("*") and not linea.startswith("**")
                if cursiva:
                    texto_en_linea(p, linea.strip()[1:-1], 9)
                    for r in p.runs:
                        r.italic = True
                        r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
                else:
                    texto_en_linea(p, linea)
            i += 1


def reiniciar_numeracion(doc):
    """Word continúa la numeración de "List Number" en todo el documento.
    Se crea una lista nueva cada vez que empieza un bloque numerado."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    numbering = doc.part.numbering_part.element
    estilo = doc.styles["List Number"]
    num_id_estilo = estilo.element.pPr.numPr.numId.val
    abstract = None
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) == str(num_id_estilo):
            abstract = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    anterior_numerado = False
    siguiente_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
    actual = None
    for p in doc.paragraphs:
        es_numerado = p.style.name == "List Number"
        if es_numerado and not anterior_numerado:
            num = OxmlElement("w:num")
            num.set(qn("w:numId"), str(siguiente_id))
            abs_el = OxmlElement("w:abstractNumId")
            abs_el.set(qn("w:val"), abstract)
            num.append(abs_el)
            override = OxmlElement("w:lvlOverride")
            override.set(qn("w:ilvl"), "0")
            inicio = OxmlElement("w:startOverride")
            inicio.set(qn("w:val"), "1")
            override.append(inicio)
            num.append(override)
            numbering.append(num)
            actual = siguiente_id
            siguiente_id += 1
        if es_numerado:
            p_pr = p._p.get_or_add_pPr()
            num_pr = OxmlElement("w:numPr")
            ilvl = OxmlElement("w:ilvl")
            ilvl.set(qn("w:val"), "0")
            num_id = OxmlElement("w:numId")
            num_id.set(qn("w:val"), str(actual))
            num_pr.append(ilvl)
            num_pr.append(num_id)
            p_pr.append(num_pr)
        anterior_numerado = es_numerado


def main():
    renderizar = "--renderizar" in sys.argv
    paginas = paginas_implementacion()

    doc = ew.nuevo_documento()
    ew.portada(
        doc,
        "Plan de implementación de AWP",
        "Advanced Work Packaging aplicado a un proyecto de construcción en tres fases",
        [
            ("Proyecto", ew.PROYECTO),
            ("Documento", "PT-AWP-PLN-001 – Plan de implementación de AWP"),
            ("Revisión", "0 – Emitido para uso"),
            ("Fecha", ew.fecha_hoy()),
            ("Elaborado por", "AWP Champion"),
            ("Revisado por", "Gerente de Proyecto"),
            ("Aprobado por", "Gerente del proyecto del cliente"),
        ],
        "Este documento se genera a partir de la sección «Implementación» de la web "
        "Consulta AWP; ambos formatos tienen el mismo contenido.",
    )

    # Índice previo (Word lo reemplaza con números de página al actualizar).
    previo = []
    for n, (nombre, ruta) in enumerate(paginas, 1):
        texto = (DOCS / ruta).read_text(encoding="utf-8")
        previo.append((1, f"{n}.  {TITULOS.get(ruta, nombre)}"))
        sub = 0
        en_codigo = False
        for l in texto.splitlines():
            if l.startswith("```"):
                en_codigo = not en_codigo
            if not en_codigo and l.startswith("## "):
                sub += 1
                previo.append((2, f"{n}.{sub}  {quitar_numero(l[3:])}"))
    ew.indice(doc, previo)

    conversor = Conversor(doc, renderizar)
    for n, (nombre, ruta) in enumerate(paginas):
        if n > 0:
            ew.salto_pagina(doc)
        texto = (DOCS / ruta).read_text(encoding="utf-8")
        conversor.pagina(texto, TITULOS.get(ruta))

    reiniciar_numeracion(doc)
    ew.encabezado_pie(doc, "Plan de implementación de AWP", "PT-AWP-PLN-001 Rev. 0")
    doc.core_properties.title = "Plan de implementación de AWP"
    doc.core_properties.subject = ew.PROYECTO
    doc.core_properties.author = "AWP Champion"
    doc.core_properties.language = "es-ES"
    doc.save(SALIDA)
    print(f"Generado: {SALIDA.relative_to(RAIZ)} ({conversor.figura} figuras)")


if __name__ == "__main__":
    main()
