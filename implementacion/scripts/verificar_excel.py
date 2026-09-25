"""Verifica las plantillas Excel recalculadas por LibreOffice.

Uso:
    soffice --headless --calc --convert-to xlsx --outdir <carpeta> implementacion/plantillas/*.xlsx
    python implementacion/scripts/verificar_excel.py <carpeta> [--detalle]

Comprueba que ninguna celda con fórmula dé error, que las listas desplegables
apunten a rangos con nombre existentes y muestra los valores calculados de
las filas de ejemplo.
"""

import sys
from pathlib import Path

from openpyxl import load_workbook

ERRORES = ("#VALUE!", "#REF!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:")
ORIGINALES = Path(__file__).resolve().parents[1] / "plantillas"


def verificar(recalculado, detalle):
    nombre = recalculado.name
    original = load_workbook(ORIGINALES / nombre)
    valores = load_workbook(recalculado, data_only=True)
    problemas = 0
    n_formulas = 0
    for ws in original.worksheets:
        wv = valores[ws.title]
        for fila in ws.iter_rows():
            for celda in fila:
                if isinstance(celda.value, str) and celda.value.startswith("="):
                    n_formulas += 1
                    v = wv[celda.coordinate].value
                    if isinstance(v, str) and v.startswith(ERRORES):
                        problemas += 1
                        if problemas <= 10:
                            print(f"  ERROR {ws.title}!{celda.coordinate}: {v}  ← {celda.value[:120]}")
        # Listas desplegables
        for dv in ws.data_validations.dataValidation:
            f = dv.formula1 or ""
            if f.startswith("=") and f[1:] not in original.defined_names:
                problemas += 1
                print(f"  LISTA sin rango con nombre: {ws.title} {f}")
        if detalle and ws.title not in ("Instrucciones", "Listas"):
            print(f"  --- {ws.title}")
            for fila in wv.iter_rows(min_row=3, max_row=18, values_only=True):
                vals = [v for v in fila if v not in (None, "")]
                if vals:
                    print("   ", [str(v)[:22] for v in vals][:26])
    n_dv = sum(len(ws.data_validations.dataValidation) for ws in original.worksheets)
    print(f"{nombre}: {n_formulas} fórmulas, {n_dv} listas desplegables, {problemas} problemas")
    return problemas


def main():
    carpeta = Path(sys.argv[1])
    detalle = "--detalle" in sys.argv
    total = sum(verificar(f, detalle) for f in sorted(carpeta.glob("*.xlsx")))
    print("TOTAL DE PROBLEMAS:", total)
    sys.exit(1 if total else 0)


if __name__ == "__main__":
    main()
