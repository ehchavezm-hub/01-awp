"""Genera las plantillas Excel del kit de implementación AWP en
implementacion/plantillas/.

Uso:
    python implementacion/scripts/generar_excel.py
"""

import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
import datos_proyecto as dp  # noqa: E402
from excel_base import (  # noqa: E402
    AMBAR, AZUL, AZUL_CLARO, BLANCO, BORDE, FILA_DATOS, FILA_ENC, FILA_PARAM, FUENTE, GRIS,
    ROJO, ROJO_TEXTO, VERDE, VERDE_TEXTO, Col, Libro, configurar_impresion, estilo_titulo,
    hoja_datos, hoja_instrucciones, relleno, resaltar_fila, semaforo_texto,
)

SALIDA = Path(__file__).resolve().parents[1] / "plantillas"
D = date
EJ = "EJEMPLO"
COL_EJEMPLO = Col("ejemplo", "Ejemplo", 10, desc="Contiene «EJEMPLO» en las filas de ejemplo. Déjela vacía en las filas reales.")


def rango(c, clave, ultima):
    return f"${c(clave)}${FILA_DATOS}:${c(clave)}${ultima}"


def tabla_resumen(ws, fila, col, titulo, etiquetas, columnas, ancho_etiqueta=None):
    """Tabla de resumen: una fila por etiqueta y una columna por fórmula.

    columnas: lista de (titulo, funcion(celda_etiqueta) -> formula, formato)
    """
    ws.cell(row=fila, column=col, value=titulo).font = Font(name=FUENTE, bold=True, size=12, color=AZUL)
    fila += 1
    enc = [""] + [t for t, _, _ in columnas]
    for i, t in enumerate(enc):
        celda = ws.cell(row=fila, column=col + i, value=t)
        celda.font = Font(name=FUENTE, bold=True, color=BLANCO, size=10)
        celda.fill = relleno(AZUL)
        celda.border = BORDE
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila].height = 32
    fila += 1
    primera = fila
    for etiqueta in etiquetas:
        celda = ws.cell(row=fila, column=col, value=etiqueta)
        celda.font = Font(name=FUENTE, bold=True, size=10)
        celda.border = BORDE
        celda.fill = relleno(AZUL_CLARO)
        ref = f"${get_column_letter(col)}{fila}"
        for i, (_, f, fmt) in enumerate(columnas, start=1):
            c = ws.cell(row=fila, column=col + i, value=f(ref))
            c.border = BORDE
            c.font = Font(name=FUENTE, size=10)
            c.alignment = Alignment(horizontal="center")
            if fmt:
                c.number_format = fmt
        fila += 1
    if ancho_etiqueta:
        ws.column_dimensions[get_column_letter(col)].width = ancho_etiqueta
    return primera, fila


# ==========================================================================
# Datos de ejemplo compartidos
# ==========================================================================

RESTRICCIONES = [
    # id, fase, cwp, iwp, tipo, descripción, rol, nombre, identificada, requerida, liberación, estado, obs
    ("R-0001", "Fase 1", "CWP-1.01-CIV-01", "IWP-1.01-CIV-01-001", "Permisos", "Permiso de excavación de la plataforma norte", "Líder de HSE", "(nombre)", D(2027, 5, 10), D(2027, 5, 31), D(2027, 6, 7), "Liberada", "Se identificó tarde; ver lección LA-03."),
    ("R-0002", "Fase 1", "CWP-1.02-TUB-01", "IWP-1.02-TUB-01-001", "Materiales", "Tubería HDPE DN200 del tramo T1 no recibida", "Líder de Procura", "(nombre)", D(2027, 7, 5), D(2027, 7, 19), D(2027, 7, 16), "Liberada", ""),
    ("R-0003", "Fase 1", "CWP-1.02-ELE-01", "IWP-1.02-ELE-01-001", "Trabajos predecesores", "Excavación de zanja para bancos BD-01 a BD-04 no terminada", "Superintendente / capataz general", "(nombre)", D(2027, 7, 19), D(2027, 8, 9), D(2027, 8, 6), "Liberada", ""),
    ("R-0004", "Fase 2", "CWP-2.01-CIV-01", "IWP-2.01-CIV-01-001", "Ingeniería", "RFI-045: detalle de pernos de anclaje de columnas A1–A4", "Líder de Ingeniería", "(nombre)", D(2027, 12, 6), D(2028, 1, 17), D(2028, 1, 24), "Liberada", ""),
    ("R-0005", "Fase 2", "CWP-2.01-EST-01", "IWP-2.01-EST-01-001", "Equipos de construcción", "Grúa de 100 t no disponible para montaje de columnas", "Gerente de Construcción", "(nombre)", D(2028, 2, 20), D(2028, 4, 11), None, "En gestión", "Escalada al Gerente de Proyecto."),
    ("R-0006", "Fase 2", "CWP-2.01-TUB-01", "IWP-2.01-TUB-01-001", "Documentación del proveedor", "Planos certificados de válvulas con tag pendientes", "Líder de Procura", "(nombre)", D(2028, 5, 2), D(2028, 6, 10), None, "Abierta", ""),
    ("R-0007", "Fase 2", "CWP-2.03-ELE-01", "IWP-2.03-ELE-01-001", "Andamios", "Andamio para montaje de bandejas en sala eléctrica", "Superintendente / capataz general", "(nombre)", D(2028, 6, 1), D(2028, 6, 20), None, "En gestión", ""),
    ("R-0008", "Fase 2", "CWP-2.01-EST-01", "IWP-2.01-EST-01-001", "HSE", "Plan de izaje crítico no aprobado", "Líder de HSE", "(nombre)", D(2028, 3, 1), D(2028, 4, 11), None, "Abierta", ""),
    ("R-0009", "Fase 3", "CWP-3.02-TUB-01", "IWP-3.02-TUB-01-001", "Interfaz entre fases", "Parada de la línea de agua de la galería 2.04 para empalme", "AWP Champion", "(nombre)", D(2028, 6, 5), D(2029, 6, 18), None, "Abierta", "Coordinar con comisionamiento de la Fase 2."),
    ("R-0010", "Fase 3", "CWP-3.01-ARQ-01", "IWP-3.01-ARQ-01-001", "Ingeniería", "EWP-3.01-ARQ-01 pendiente de emisión IFC", "Líder de Ingeniería", "(nombre)", D(2028, 6, 12), D(2029, 1, 14), None, "Abierta", ""),
]
CLAVES_RESTR = ["id", "fase", "cwp", "iwp", "tipo", "desc", "rol", "nombre", "f_ident", "f_req", "f_lib", "estado", "obs"]


def ejemplos(tuplas, claves):
    return [dict({"ejemplo": EJ}, **dict(zip(claves, t))) for t in tuplas]


# ==========================================================================
# 1. Registro de restricciones
# ==========================================================================

def registro_restricciones():
    lib = Libro(SALIDA / "Registro_Restricciones.xlsx", "Registro de restricciones",
                ["Fases", "TiposRestriccion", "EstadosRestriccion", "Roles"])

    def dias(f, c, p):
        return (f'=IF(OR({c("id")}{f}="",{c("f_req")}{f}=""),"",IF({c("estado")}{f}="Cancelada",0,'
                f'IF({c("estado")}{f}="Liberada",IF({c("f_lib")}{f}="",0,MAX(0,{c("f_lib")}{f}-{c("f_req")}{f})),'
                f'MAX(0,{p["corte"]}-{c("f_req")}{f}))))')

    def situacion(f, c, p):
        e, r = f'{c("estado")}{f}', f'{c("f_req")}{f}'
        return (f'=IF({c("id")}{f}="","",IF({e}="Cancelada","Cancelada",IF({e}="Liberada",'
                f'IF({c("dias")}{f}>0,"Liberada con atraso","Liberada a tiempo"),IF({r}="","Sin fecha requerida",'
                f'IF({p["corte"]}>{r},"Vencida",IF({r}-{p["corte"]}<=7,"Por vencer","En plazo"))))))')

    cols = [
        COL_EJEMPLO,
        Col("id", "ID", 9, desc="Identificador único: R-0001, R-0002… No se reutiliza aunque la restricción se cancele."),
        Col("fase", "Fase del proyecto", 11, lista="Fases", desc="Fase del proyecto a la que pertenece el IWP o CWP afectado."),
        Col("cwp", "CWP relacionado", 17, desc="Código del CWP (formato CWP-<fase>.<nn>-<disciplina>-<nn>)."),
        Col("iwp", "IWP relacionado", 21, desc="Código del IWP afectado (formato IWP-<fase>.<nn>-<disciplina>-<nn>-<nnn>). Vacío si la restricción afecta a todo el CWP."),
        Col("tipo", "Tipo de restricción", 18, lista="TiposRestriccion", desc="Categoría de la restricción (lista común del kit)."),
        Col("desc", "Descripción", 40, ajustar=True, desc="Qué falta exactamente para poder ejecutar el trabajo. Debe ser verificable."),
        Col("rol", "Responsable (rol)", 22, lista="Roles", desc="Rol responsable de levantar la restricción."),
        Col("nombre", "Nombre del responsable", 16, desc="Persona concreta que levantará la restricción."),
        Col("f_ident", "Fecha identificada", 12, "fecha", desc="Fecha en que el planificador registró la restricción."),
        Col("f_req", "Fecha requerida", 12, "fecha", desc="Fecha límite para liberar la restricción sin afectar la liberación del IWP."),
        Col("f_lib", "Fecha de liberación", 12, "fecha", desc="Fecha real en que se verificó la liberación. Solo si el estado es «Liberada»."),
        Col("estado", "Estado", 12, lista="EstadosRestriccion", desc="Abierta: registrada sin gestión. En gestión: con responsable trabajando en ella. Liberada: resuelta y verificada. Cancelada: ya no aplica."),
        Col("dias", "Días de atraso", 10, "num", formula=dias, desc="Calculado. Abiertas/en gestión: días desde la fecha requerida hasta la fecha de corte. Liberadas: días entre la fecha requerida y la de liberación. 0 si no hay atraso."),
        Col("situacion", "Situación", 17, formula=situacion, desc="Calculado: Vencida, Por vencer (≤ 7 días), En plazo, Liberada a tiempo, Liberada con atraso o Cancelada."),
        Col("obs", "Observaciones", 34, ajustar=True, desc="Comentarios, escalamientos, evidencia de liberación."),
    ]
    ws, c, ult = hoja_datos(lib, "Registro", "Registro de restricciones", cols,
                            ejemplos(RESTRICCIONES, CLAVES_RESTR), n_filas=300,
                            params=[("corte", "Fecha de corte:", dp.FECHA_CORTE, "fecha")], congelar_col=2)
    todo = f"A{FILA_DATOS}:{c('obs')}{ult}"
    s = f"${c('situacion')}{FILA_DATOS}"
    resaltar_fila(ws, todo, f'{s}="Vencida"', ROJO, ROJO_TEXTO)
    resaltar_fila(ws, todo, f'{s}="Por vencer"', AMBAR)
    resaltar_fila(ws, f"{c('situacion')}{FILA_DATOS}:{c('situacion')}{ult}", f'OR({s}="Liberada a tiempo",{s}="Liberada con atraso")', VERDE, VERDE_TEXTO)

    # Resumen
    rs = lib.wb.create_sheet("Resumen")
    estilo_titulo(rs, "Resumen de restricciones", f"{dp.PROYECTO}. Se calcula automáticamente a partir de la hoja «Registro».", 8)
    rf, re_, rt, rsit, rd = (f"Registro!{rango(c, k, ult)}" for k in ("fase", "estado", "tipo", "situacion", "dias"))
    cols_estado = [(e, (lambda ref, e=e: f'=COUNTIFS({rf},{ref},{re_},"{e}")'), "0") for e in dp.ESTADOS_RESTRICCION]
    cols_estado += [
        ("Total", lambda ref: f'=COUNTIFS({rf},{ref},Registro!{rango(c, "id", ult)},"<>")', "0"),
        ("Vencidas", lambda ref: f'=COUNTIFS({rf},{ref},{rsit},"Vencida")', "0"),
        ("Por vencer (≤ 7 días)", lambda ref: f'=COUNTIFS({rf},{ref},{rsit},"Por vencer")', "0"),
        ("Atraso promedio de vencidas (días)", lambda ref: f'=IFERROR(AVERAGEIFS({rd},{rf},{ref},{rsit},"Vencida"),0)', "0.0"),
        ("% liberadas a tiempo", lambda ref: f'=IFERROR(COUNTIFS({rf},{ref},{rsit},"Liberada a tiempo")/COUNTIFS({rf},{ref},{re_},"Liberada"),"")', "0%"),
    ]
    tabla_resumen(rs, 4, 1, "Por fase del proyecto", dp.FASES, cols_estado, 22)
    cols_tipo = [
        ("Abiertas + en gestión", lambda ref: f'=COUNTIFS({rt},{ref},{re_},"Abierta")+COUNTIFS({rt},{ref},{re_},"En gestión")', "0"),
        ("Vencidas", lambda ref: f'=COUNTIFS({rt},{ref},{rsit},"Vencida")', "0"),
        ("Liberadas", lambda ref: f'=COUNTIFS({rt},{ref},{re_},"Liberada")', "0"),
    ] + [(f, (lambda ref, f=f: f'=COUNTIFS({rt},{ref},{rf},"{f}")'), "0") for f in dp.FASES]
    tabla_resumen(rs, 11, 1, "Por tipo de restricción", dp.TIPOS_RESTRICCION, cols_tipo, 26)
    for i in range(2, 11):
        rs.column_dimensions[get_column_letter(i)].width = 14
    configurar_impresion(rs, "Resumen de restricciones", area="A1:J26")

    hoja_instrucciones(
        lib,
        "Registrar en una sola lista todas las restricciones que impiden liberar un IWP o iniciar un CWP, con su responsable y fecha requerida, para levantarlas antes de que el trabajo llegue a campo. Es la herramienta central de la «regla de oro» de WFP: ningún IWP se libera con restricciones abiertas.",
        "Cada semana, desde que el IWP entra a la ventana de planificación (8 a 12 semanas antes de su ejecución según la fase) hasta que se libera. Se revisa en la reunión semanal de restricciones.",
        "Los planificadores de frente de trabajo registran las restricciones; cada responsable actualiza el estado y la fecha de liberación; el Líder de WFP consolida y revisa la calidad del registro.",
        [
            "Actualice la fecha de corte (celda B3) con la fecha del reporte. Para que se actualice sola, escriba =HOY().",
            "Borre las filas de ejemplo (amarillas) antes de usar la plantilla en un proyecto real.",
            "Registre una fila por restricción. Use los códigos de CWP e IWP de la plantilla Seguimiento_Paquetes.",
            "Elija fase, tipo, responsable y estado de las listas desplegables.",
            "Fije la fecha requerida según los plazos de la fase (ver Programa_Liberacion_IWP).",
            "Cuando la restricción se resuelva, cambie el estado a «Liberada» y anote la fecha de liberación y la evidencia en observaciones.",
            "No escriba en las columnas calculadas (encabezado azul medio y celdas grises).",
            "Use los filtros del encabezado para preparar la reunión semanal (por ejemplo, Situación = Vencida o Por vencer).",
            "Revise la hoja «Resumen» para el reporte semanal por fase.",
        ],
        [("Registro", cols)],
        notas=[("Formato condicional", "Fila en rojo: restricción vencida (no liberada y con fecha requerida anterior a la fecha de corte). Fila en ámbar: vence en 7 días o menos. Situación en verde: liberada.")],
    )
    lib.crear_listas()
    lib.guardar()


# ==========================================================================
# 2. Matriz RACI
# ==========================================================================

ACTIVIDADES_RACI = [
    ("Planificación temprana (FEL)", "Aprobar la estrategia y el plan de implementación AWP", "Todas las fases", dict(CLI="A", GP="R", CHA="R", CON="C", ING="C", PRO="C", CTR="C", COM="I")),
    ("Planificación temprana (FEL)", "Designar al AWP Champion y al equipo AWP", "Todas las fases", dict(CLI="C", GP="A", CHA="I", CON="C")),
    ("Planificación temprana (FEL)", "Evaluar la escala de AWP y definir el nivel de madurez por fase", "Todas las fases", dict(CLI="C", GP="A", CHA="R", CON="C", CTR="C")),
    ("Planificación temprana (FEL)", "Definir las CWA de la fase", "Todas las fases", dict(CLI="C", CHA="A", CON="R", ING="C", COM="C", CTR="C", IM="I")),
    ("Planificación temprana (FEL)", "Desarrollar el Path of Construction en talleres IPP", "Todas las fases", dict(CLI="C", CHA="R", CON="A", ING="C", PRO="C", COM="C", CTR="C", HSE="C", LWF="C")),
    ("Planificación temprana (FEL)", "Elaborar el plan de liberación CWP/EWP/PWP", "Todas las fases", dict(CHA="A", CON="R", ING="R", PRO="R", CTR="R", LWF="C")),
    ("Planificación temprana (FEL)", "Definir la WBS y la codificación de paquetes", "Todas las fases", dict(CHA="A", IM="R", CTR="R", ING="C", PRO="C", CON="C")),
    ("Planificación temprana (FEL)", "Incluir los requisitos AWP en contratos y subcontratos", "Todas las fases", dict(CLI="C", GP="A", CHA="C", PRO="R", SUB="I")),
    ("Planificación temprana (FEL)", "Definir los requisitos de gestión de la información y del modelo 3D", "Todas las fases", dict(CHA="A", IM="R", ING="C", CTR="C", PRO="C")),
    ("Planificación temprana (FEL)", "Implementar el software de WFP integrado con modelo, cronograma y materiales", "Fase 2", dict(CHA="A", IM="R", LWF="C", ING="C", PRO="C", CTR="C")),
    ("Ingeniería", "Planificar la ingeniería por EWP según el PoC", "Todas las fases", dict(ING="A", CTR="R", CON="C", CHA="C")),
    ("Ingeniería", "Emitir EWP completos (IFC y MTO) en la secuencia del PoC", "Todas las fases", dict(ING="A", CON="I", PRO="I", LWF="I", CHA="I")),
    ("Ingeniería", "Mantener el modelo 3D con atributos AWP", "Todas las fases", dict(IM="A", ING="R", LWF="C", WFP="I")),
    ("Ingeniería", "Revisar la constructabilidad antes de la emisión IFC", "Todas las fases", dict(ING="A", CON="R", SUP="C", HSE="C", CAL="C")),
    ("Ingeniería", "Responder las RFI", "Todas las fases", dict(ING="A", WFP="C", LWF="I", CON="I")),
    ("Procura", "Estructurar las compras en PWP por CWP", "Todas las fases", dict(PRO="A", ING="C", CON="C", MAT="C", CHA="C")),
    ("Procura", "Alinear las fechas RAS con el PoC", "Todas las fases", dict(PRO="A", CTR="R", CON="C", ING="C")),
    ("Procura", "Informar el estado de materiales por CWP e IWP", "Todas las fases", dict(PRO="A", MAT="R", LWF="I", WFP="I")),
    ("Procura", "Reservar y despachar materiales por IWP", "Todas las fases", dict(MAT="A", WFP="C", SUP="I", PRO="C")),
    ("Procura", "Gestionar la documentación del proveedor", "Todas las fases", dict(PRO="A", ING="C", CAL="C", COM="C")),
    ("Construcción", "Emitir los CWP", "Todas las fases", dict(CON="A", LWF="R", ING="C", PRO="C", HSE="C", CAL="C", CTR="C")),
    ("Construcción", "Dividir los CWP en IWP", "Todas las fases", dict(LWF="A", WFP="R", SUP="C", CTR="I")),
    ("Construcción", "Identificar y registrar restricciones", "Todas las fases", dict(LWF="A", WFP="R", SUP="C", SUB="R")),
    ("Construcción", "Levantar las restricciones asignadas", "Todas las fases", dict(LWF="A", ING="R", PRO="R", MAT="R", HSE="R", SUP="R")),
    ("Construcción", "Dirigir la reunión semanal de restricciones", "Todas las fases", dict(LWF="A", CON="C", WFP="C", ING="C", PRO="C", SUP="C")),
    ("Construcción", "Liberar IWP con el checklist de liberación", "Todas las fases", dict(LWF="A", WFP="R", SUP="C", HSE="C", CAL="C", MAT="C")),
    ("Construcción", "Seleccionar IWP del backlog (lookahead de 3 semanas)", "Todas las fases", dict(CON="A", SUP="R", LWF="C", SUB="C")),
    ("Construcción", "Ejecutar el IWP en campo", "Todas las fases", dict(SUP="A", SUB="R", WFP="I", CAL="C", HSE="C")),
    ("Construcción", "Cerrar el IWP (avance, HH, calidad)", "Todas las fases", dict(LWF="A", WFP="R", SUP="R", CAL="R", CTR="I")),
    ("Construcción", "Calcular y reportar los KPI", "Todas las fases", dict(CHA="A", CTR="R", LWF="R", GP="I", CLI="I")),
    ("Construcción", "Gestionar las interfaces entre fases superpuestas", "Fases 2 y 3", dict(CHA="A", CON="R", ING="C", COM="C", CTR="C")),
    ("Construcción", "Auditar el proceso AWP", "Todas las fases", dict(CHA="A", LWF="C", CON="C", GP="I", CLI="I")),
    ("Comisionamiento", "Definir los sistemas y la secuencia de puesta en marcha", "Todas las fases", dict(COM="A", ING="R", CON="C", CHA="C", CLI="C")),
    ("Comisionamiento", "Relacionar los IWP con los sistemas", "Todas las fases", dict(COM="A", LWF="R", IM="R", WFP="C")),
    ("Comisionamiento", "Preparar los SWP y TOP", "Todas las fases", dict(COM="A", CAL="R", CON="C", SUP="C", CLI="I")),
    ("Comisionamiento", "Capturar y transferir lecciones aprendidas", "Todas las fases", dict(CHA="A", LWF="R", GP="C", CLI="I", ING="C", PRO="C", CON="C", COM="C")),
    ("Comisionamiento", "Elaborar el informe de cierre AWP de la fase", "Todas las fases", dict(CHA="A", CTR="C", LWF="C", GP="C", CLI="I")),
]


def matriz_raci():
    lib = Libro(SALIDA / "Matriz_Roles_RACI.xlsx", "Matriz de roles RACI", ["Etapas", "FasesExt", "Roles", "CodigosRol"])
    codigos = [c for c, _ in dp.ROLES]

    def cuenta(letra):
        def f(fila, c, p):
            return f'=IF({c("act")}{fila}="","",COUNTIF({c(codigos[0])}{fila}:{c(codigos[-1])}{fila},"{letra}"))'
        return f

    def verif(fila, c, p):
        a, r = f'{c("na")}{fila}', f'{c("nr")}{fila}'
        return f'=IF({c("act")}{fila}="","",IF({a}<>1,"Revisar: debe haber una sola A",IF({r}=0,"OK (A ejecuta)","OK")))'

    cols = [
        Col("n", "N.°", 5, "num", desc="Número correlativo de la actividad."),
        Col("etapa", "Etapa del ciclo de vida", 20, lista="Etapas", desc="Etapa del ciclo de vida en que se realiza la actividad."),
        Col("act", "Actividad AWP", 46, ajustar=True, desc="Actividad del proceso AWP."),
        Col("fase", "Fase del proyecto", 13, lista="FasesExt", desc="Fases en que aplica la actividad."),
    ]
    for cod, nombre in dp.ROLES:
        cols.append(Col(cod, cod, 5.5, lista_valores=["R", "A", "C", "I"], desc=f"{nombre}. R, A, C o I (vacío si no participa)."))
    cols += [
        Col("na", "N.° de A", 7, "num", formula=cuenta("A"), desc="Calculado: cantidad de «A» en la fila. Debe ser exactamente 1."),
        Col("nr", "N.° de R", 7, "num", formula=cuenta("R"), desc="Calculado: cantidad de «R» en la fila."),
        Col("ver", "Verificación", 24, formula=verif, desc="Calculado: «OK» si hay una sola A; «Revisar» si no."),
    ]
    filas = []
    for i, (etapa, act, fase, asign) in enumerate(ACTIVIDADES_RACI, 1):
        filas.append(dict({"n": i, "etapa": etapa, "act": act, "fase": fase}, **asign))
    ws, c, ult = hoja_datos(
        lib, "Matriz RACI", "Matriz de roles y responsabilidades AWP (RACI)", cols, [], n_filas=60,
        subtitulo=f"{dp.PROYECTO}. Matriz propuesta para el proyecto tipo: ajústela en el taller de arranque de cada fase. R = responsable de ejecutar; A = aprueba y rinde cuentas (una sola por actividad); C = consultado; I = informado.",
        congelar_col=4, alto_enc=30,
    )
    # La matriz propuesta se escribe sin fondo amarillo: es la base del proyecto.
    for n, datos in enumerate(filas):
        fila = FILA_DATOS + n
        for col in cols:
            if col.clave in datos:
                ws[f"{c(col.clave)}{fila}"] = datos[col.clave]
    for cod in codigos:
        ws[f"{c(cod)}{FILA_ENC}"].alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)
    ws.row_dimensions[FILA_ENC].height = 40
    rr = f"{c(codigos[0])}{FILA_DATOS}:{c(codigos[-1])}{ult}"
    p = f"{c(codigos[0])}{FILA_DATOS}"
    resaltar_fila(ws, rr, f'{p}="A"', "1F3864", BLANCO, True)
    resaltar_fila(ws, rr, f'{p}="R"', "9DC3E6", "000000", True)
    resaltar_fila(ws, rr, f'{p}="C"', "E2EFDA")
    resaltar_fila(ws, rr, f'{p}="I"', "EDEDED")
    for cod in codigos:
        for fila in range(FILA_DATOS, ult + 1):
            ws[f"{c(cod)}{fila}"].alignment = Alignment(horizontal="center")
    v = f"${c('ver')}{FILA_DATOS}"
    resaltar_fila(ws, f"{c('ver')}{FILA_DATOS}:{c('ver')}{ult}", f'LEFT({v},7)="Revisar"', ROJO, ROJO_TEXTO)
    resaltar_fila(ws, f"{c('ver')}{FILA_DATOS}:{c('ver')}{ult}", f'LEFT({v},2)="OK"', VERDE, VERDE_TEXTO)

    # Hoja de roles
    rs = lib.wb.create_sheet("Roles")
    estilo_titulo(rs, "Roles AWP del proyecto tipo", "Códigos usados en todas las plantillas del kit.", 3)
    responsabilidades = {
        "CLI": "Patrocina AWP; aprueba el plan, el PoC y los hitos; exige requisitos AWP en contratos.",
        "GP": "Responsable final de la implementación; asigna recursos; resuelve restricciones escaladas.",
        "CHA": "Lidera la implementación; procedimientos, plantillas, formación, talleres de PoC, auditorías, KPI y lecciones.",
        "LWF": "Coordina a los planificadores; dirige la reunión de restricciones; controla backlog y liberación de IWP.",
        "WFP": "Divide CWP en IWP; identifica restricciones; arma, libera y cierra IWP (1 por cada 50 trabajadores).",
        "ING": "Emite EWP en secuencia del PoC; mantiene el modelo 3D; responde RFI.",
        "PRO": "Estructura PWP; alinea fechas RAS; informa estado de materiales; documentación del proveedor.",
        "MAT": "Recibe, reserva y despacha materiales por IWP.",
        "CON": "Lidera el PoC; aprueba CWP; decide la secuencia de ejecución; prioriza restricciones críticas.",
        "SUP": "Selecciona IWP del backlog; ejecuta los IWP con sus capataces; informa avance.",
        "CTR": "Integra CWP (nivel 3) e IWP (nivel 5) en el cronograma; mide avance y productividad; calcula KPI.",
        "IM": "Codificación y atributos; integración de modelo, cronograma, materiales y documentos; software de WFP.",
        "HSE": "JHA y permisos en CWP e IWP; valida requisitos de seguridad antes de liberar.",
        "CAL": "ITP y registros de inspección en CWP e IWP; cierre de calidad de cada IWP.",
        "COM": "Sistemas y secuencia de puesta en marcha; relación IWP–sistema; SWP y TOP.",
        "SUB": "Aplica AWP en su alcance; aporta planificadores; cumple el lookahead.",
    }
    for i, t in enumerate(("Código", "Rol", "Responsabilidades AWP principales"), 1):
        celda = rs.cell(row=4, column=i, value=t)
        celda.font = Font(name=FUENTE, bold=True, color=BLANCO)
        celda.fill = relleno(AZUL)
        celda.border = BORDE
    for n, (cod, nombre) in enumerate(dp.ROLES, start=5):
        for i, v in enumerate((cod, nombre, responsabilidades[cod]), 1):
            celda = rs.cell(row=n, column=i, value=v)
            celda.border = BORDE
            celda.alignment = Alignment(wrap_text=True, vertical="top")
            celda.font = Font(name=FUENTE, size=10, bold=(i == 1))
    rs.column_dimensions["A"].width = 9
    rs.column_dimensions["B"].width = 34
    rs.column_dimensions["C"].width = 80
    rs.freeze_panes = "A5"
    configurar_impresion(rs, "Roles AWP", "landscape", "4:4", f"A1:C{4 + len(dp.ROLES)}")

    # Resumen de carga por rol
    rr_ = lib.wb.create_sheet("Resumen por rol")
    estilo_titulo(rr_, "Participación de cada rol", "Cantidad de actividades en que cada rol es R, A, C o I (calculado).", 6)
    fila_enc = 4
    for i, t in enumerate(("Código", "Rol", "R", "A", "C", "I"), 1):
        celda = rr_.cell(row=fila_enc, column=i, value=t)
        celda.font = Font(name=FUENTE, bold=True, color=BLANCO)
        celda.fill = relleno(AZUL)
        celda.border = BORDE
        celda.alignment = Alignment(horizontal="center")
    for n, (cod, nombre) in enumerate(dp.ROLES, start=fila_enc + 1):
        rr_.cell(row=n, column=1, value=cod).font = Font(name=FUENTE, bold=True)
        rr_.cell(row=n, column=2, value=nombre)
        for k, letra in enumerate("RACI", start=3):
            rr_.cell(row=n, column=k, value=f"=COUNTIF('Matriz RACI'!{rango(c, cod, ult)},\"{letra}\")")
        for k in range(1, 7):
            rr_.cell(row=n, column=k).border = BORDE
            rr_.cell(row=n, column=k).alignment = Alignment(horizontal="center" if k > 2 else None)
    rr_.column_dimensions["B"].width = 36
    configurar_impresion(rr_, "Participación por rol", "portrait", area=f"A1:F{fila_enc + len(dp.ROLES)}")

    hoja_instrucciones(
        lib,
        "Definir quién ejecuta (R), quién aprueba (A), a quién se consulta (C) y a quién se informa (I) en cada actividad AWP, por etapa del ciclo de vida, para que no queden actividades sin dueño ni con dueños duplicados.",
        "En la planificación temprana global (antes del hito H0) y en el arranque de cada fase del proyecto. Se revisa cuando cambia la organización o se incorpora un subcontratista.",
        "El AWP Champion la prepara con los líderes funcionales; la aprueba el Gerente de Proyecto y se comunica a todo el equipo.",
        [
            "Revise la lista de actividades y agregue las propias del proyecto al final (la numeración es libre).",
            "Para cada actividad, elija R, A, C o I en la columna de cada rol (lista desplegable). Deje vacío si el rol no participa.",
            "Asigne exactamente una «A» por actividad. La columna «Verificación» avisa en rojo si hay cero o más de una.",
            "Indique en «Fase del proyecto» si la actividad aplica a todas las fases o a una en particular.",
            "Revise la hoja «Resumen por rol» para detectar roles sobrecargados o sin participación.",
            "La hoja «Roles» describe cada código de rol; los mismos códigos se usan en todo el kit.",
        ],
        [("Matriz RACI", [col for col in cols if col.clave not in [c_ for c_, _ in dp.ROLES]] +
          [("Columnas de roles (CLI … SUB)", "Una columna por rol con R, A, C o I. Ver la hoja «Roles».", "Lista")])],
        notas=[("Nota", "A diferencia de otras plantillas, esta matriz no tiene filas de ejemplo: es la propuesta del proyecto tipo y se ajusta directamente.")],
    )
    lib.crear_listas()
    lib.guardar()


# ==========================================================================
# 3. Seguimiento de paquetes
# ==========================================================================

PAQUETES = [
    # código, tipo, fase, descripción, disciplina, cwa, cwp, sistema, rol, estado, ini plan, fin plan, ini real, fin real, avance, HH, obs
    ("CWA-1.01", "CWA", "Fase 1", "Plataformas y movimiento de tierras", "CIV", "CWA-1.01", "", "Plataformas", "Gerente de Construcción", "Cerrado", D(2027, 6, 7), D(2027, 12, 20), D(2027, 6, 14), D(2028, 1, 10), 1, 120000, "Entregada a la Fase 2 con acta."),
    ("CWP-1.01-CIV-01", "CWP", "Fase 1", "Movimiento de tierras plataforma norte", "CIV", "CWA-1.01", "CWP-1.01-CIV-01", "Plataformas", "Gerente de Construcción", "Cerrado", D(2027, 6, 7), D(2027, 7, 25), D(2027, 6, 14), D(2027, 8, 6), 1, 42000, ""),
    ("EWP-1.01-CIV-01", "EWP", "Fase 1", "Ingeniería de movimiento de tierras plataforma norte", "CIV", "CWA-1.01", "CWP-1.01-CIV-01", "Plataformas", "Líder de Ingeniería", "Cerrado", D(2027, 3, 1), D(2027, 4, 30), D(2027, 3, 1), D(2027, 4, 26), 1, 0, "Emitido IFC."),
    ("PWP-1.01-CIV-01", "PWP", "Fase 1", "Material de préstamo y geotextil", "CIV", "CWA-1.01", "CWP-1.01-CIV-01", "Plataformas", "Líder de Procura", "Cerrado", D(2027, 4, 5), D(2027, 6, 1), D(2027, 4, 5), D(2027, 6, 10), 1, 0, ""),
    ("IWP-1.01-CIV-01-001", "IWP", "Fase 1", "Excavación masiva ejes 1–5", "CIV", "CWA-1.01", "CWP-1.01-CIV-01", "Plataformas", "Planificador de frente de trabajo", "Cerrado", D(2027, 6, 7), D(2027, 6, 12), D(2027, 6, 14), D(2027, 6, 19), 1, 480, "Inicio atrasado por permiso (R-0001)."),
    ("CWA-2.01", "CWA", "Fase 2", "Instalación principal – bloque A", "VAR", "CWA-2.01", "", "Varios", "Gerente de Construcción", "En ejecución", D(2028, 2, 7), D(2029, 2, 26), D(2028, 2, 12), None, 0.18, 380000, ""),
    ("CWP-2.01-EST-01", "CWP", "Fase 2", "Estructura metálica bloque A", "EST", "CWA-2.01", "CWP-2.01-EST-01", "Estructura bloque A", "Gerente de Construcción", "Liberado", D(2028, 4, 25), D(2028, 8, 21), None, None, 0, 9600, "CWP emitido; montaje no iniciado por grúa (R-0005)."),
    ("EWP-2.01-EST-01", "EWP", "Fase 2", "Ingeniería de estructura metálica bloque A", "EST", "CWA-2.01", "CWP-2.01-EST-01", "Estructura bloque A", "Líder de Ingeniería", "Cerrado", D(2027, 9, 1), D(2027, 12, 29), D(2027, 9, 1), D(2028, 1, 15), 1, 0, "IFC emitido con 17 días de atraso."),
    ("IWP-2.01-EST-01-001", "IWP", "Fase 2", "Montaje de columnas nivel 1 ejes A1–A4", "EST", "CWA-2.01", "CWP-2.01-EST-01", "Estructura bloque A", "Planificador de frente de trabajo", "En desarrollo", D(2028, 4, 25), D(2028, 4, 30), None, None, 0, 420, "No liberable: 2 restricciones abiertas."),
    ("CWP-3.02-CIV-01", "CWP", "Fase 3", "Cimentaciones de la ampliación", "CIV", "CWA-3.02", "CWP-3.02-CIV-01", "Estructura ampliación", "Gerente de Construcción", "Planificado", D(2029, 2, 5), D(2029, 5, 13), None, None, 0, 28000, ""),
]
CLAVES_PAQ = ["cod", "tipo", "fase", "desc", "disc", "cwa", "cwp", "sistema", "rol", "estado", "ini_p", "fin_p", "ini_r", "fin_r", "avance", "hh", "obs"]


def seguimiento_paquetes():
    lib = Libro(SALIDA / "Seguimiento_Paquetes.xlsx", "Seguimiento de paquetes", ["Fases", "TiposPaquete", "Disciplinas", "EstadosPaquete", "Roles", "EstadosRestriccion"])
    NR = 300  # filas de la hoja Restricciones
    fin_r = FILA_DATOS + NR - 1
    rc = f"Restricciones!$C${FILA_DATOS}:$C${fin_r}"
    ri = f"Restricciones!$D${FILA_DATOS}:$D${fin_r}"
    re_ = f"Restricciones!$F${FILA_DATOS}:$F${fin_r}"

    def abiertas(col_rango, criterio):
        return f'COUNTIFS({col_rango},{criterio},{re_},"Abierta")+COUNTIFS({col_rango},{criterio},{re_},"En gestión")'

    def restr(f, c, p):
        cod, tipo, cwp = f'{c("cod")}{f}', f'{c("tipo")}{f}', f'{c("cwp")}{f}'
        # CWA: todos los CWP de la CWA (CWP-2.01-*). IWP: el propio IWP. Resto: su CWP.
        de_cwa = '"CWP-"&MID(' + cod + ',5,4)&"-*"'
        de_cwp = f'IF({tipo}="CWP",{cod},{cwp})'
        return (f'=IF({cod}="","",IF({tipo}="IWP",{abiertas(ri, cod)},'
                f'IF({tipo}="CWA",{abiertas(rc, de_cwa)},{abiertas(rc, de_cwp)})))')

    def atraso(f, c, p):
        fp, fr, e = f'{c("fin_p")}{f}', f'{c("fin_r")}{f}', f'{c("estado")}{f}'
        return (f'=IF(OR({c("cod")}{f}="",{fp}=""),"",IF({fr}<>"",MAX(0,{fr}-{fp}),'
                f'IF(AND({p["corte"]}>{fp},{e}<>"Cerrado"),{p["corte"]}-{fp},0)))')

    def situacion(f, c, p):
        e, a, ip, ir = f'{c("estado")}{f}', f'{c("atraso")}{f}', f'{c("ini_p")}{f}', f'{c("ini_r")}{f}'
        return (f'=IF({c("cod")}{f}="","",IF({e}="Cerrado",IF({a}>0,"Cerrado con atraso","Cerrado"),'
                f'IF({a}>0,"Atrasado",IF(AND({ir}="",{ip}<>"",{p["corte"]}>{ip}),"Inicio atrasado","En plazo"))))')

    cols = [
        COL_EJEMPLO,
        Col("cod", "Código", 21, desc="Código del paquete según la codificación del kit (CWA-2.01, CWP-2.01-EST-01, EWP-…, PWP-…, IWP-…-001)."),
        Col("tipo", "Tipo", 7, lista="TiposPaquete", desc="CWA, CWP, EWP, PWP o IWP."),
        Col("fase", "Fase del proyecto", 11, lista="Fases", desc="Fase del proyecto a la que pertenece el paquete."),
        Col("desc", "Descripción", 34, ajustar=True, desc="Nombre o alcance resumido del paquete."),
        Col("disc", "Disciplina", 9, lista="Disciplinas", desc="Código de disciplina (VAR para CWA con varias disciplinas)."),
        Col("cwa", "CWA", 10, desc="CWA a la que pertenece el paquete (relación jerárquica)."),
        Col("cwp", "CWP asociado", 17, desc="CWP al que alimenta (EWP, PWP) o del que proviene (IWP). Para un CWP, su propio código. Vacío para una CWA."),
        Col("sistema", "Sistema", 16, desc="Sistema de comisionamiento al que pertenece (permite preparar los SWP)."),
        Col("rol", "Responsable (rol)", 20, lista="Roles", desc="Rol responsable del paquete."),
        Col("estado", "Estado", 12, lista="EstadosPaquete", desc="Estado del paquete (ver significado de cada estado más abajo)."),
        Col("ini_p", "Inicio plan", 11, "fecha", desc="Fecha de inicio planificada (línea base)."),
        Col("fin_p", "Fin plan", 11, "fecha", desc="Fecha de fin planificada (línea base). Para un EWP: emisión IFC; para un PWP: entrega en obra (RAS)."),
        Col("ini_r", "Inicio real", 11, "fecha", desc="Fecha de inicio real."),
        Col("fin_r", "Fin real", 11, "fecha", desc="Fecha de fin real."),
        Col("avance", "% avance", 8, "pct", desc="Avance físico del paquete (0 % a 100 %)."),
        Col("hh", "HH estimadas", 10, "num", desc="Horas-hombre directas estimadas (CWA, CWP e IWP). 0 para EWP y PWP."),
        Col("restr", "Restricciones abiertas", 11, "num", formula=restr, desc="Calculado desde la hoja «Restricciones»: restricciones abiertas o en gestión del IWP, del CWP asociado o de todos los CWP de la CWA."),
        Col("atraso", "Días de atraso", 9, "num", formula=atraso, desc="Calculado: días entre el fin plan y el fin real, o hasta la fecha de corte si el paquete no está cerrado y ya pasó su fin plan."),
        Col("situacion", "Situación", 16, formula=situacion, desc="Calculado: En plazo, Inicio atrasado, Atrasado, Cerrado o Cerrado con atraso."),
        Col("obs", "Observaciones", 32, ajustar=True, desc="Comentarios."),
    ]
    ws, c, ult = hoja_datos(lib, "Paquetes", "Seguimiento de paquetes CWA · CWP · EWP · PWP · IWP", cols,
                            ejemplos(PAQUETES, CLAVES_PAQ), n_filas=500,
                            params=[("corte", "Fecha de corte:", dp.FECHA_CORTE, "fecha")])
    s = f"${c('situacion')}{FILA_DATOS}"
    rsit = f"{c('atraso')}{FILA_DATOS}:{c('situacion')}{ult}"
    resaltar_fila(ws, rsit, f'OR({s}="Atrasado",{s}="Inicio atrasado")', ROJO, ROJO_TEXTO)
    resaltar_fila(ws, rsit, f'{s}="Cerrado con atraso"', AMBAR)
    resaltar_fila(ws, rsit, f'OR({s}="Cerrado",{s}="En plazo")', VERDE, VERDE_TEXTO)
    r = f"${c('restr')}{FILA_DATOS}"
    resaltar_fila(ws, f"{c('restr')}{FILA_DATOS}:{c('restr')}{ult}", f"AND(ISNUMBER({r}),{r}>0)", AMBAR, None, True)
    from openpyxl.formatting.rule import DataBarRule
    ws.conditional_formatting.add(f"{c('avance')}{FILA_DATOS}:{c('avance')}{ult}",
                                  DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="5B9BD5"))

    # Hoja de restricciones (copia simplificada del registro)
    rcols = [
        Col("id", "ID", 9, desc="ID de la restricción (igual que en Registro_Restricciones)."),
        Col("fase", "Fase del proyecto", 11, lista="Fases", desc="Fase del proyecto."),
        Col("cwp", "CWP relacionado", 17, desc="Código del CWP."),
        Col("iwp", "IWP relacionado", 21, desc="Código del IWP."),
        Col("tipo", "Tipo de restricción", 22, desc="Tipo de restricción."),
        Col("estado", "Estado", 12, lista="EstadosRestriccion", desc="Estado actual."),
    ]
    rej = [dict(id=t[0], fase=t[1], cwp=t[2], iwp=t[3], tipo=t[4], estado=t[11]) for t in RESTRICCIONES]
    hoja_datos(lib, "Restricciones", "Restricciones (copia del registro)", rcols, rej, n_filas=NR,
               subtitulo="Pegue aquí (solo valores) las columnas ID, Fase, CWP, IWP, Tipo y Estado de Registro_Restricciones.xlsx. Las filas amarillas son de ejemplo.",
               congelar_col=1, orientacion="portrait")
    lib.wb.move_sheet("Restricciones", offset=1)

    # Resumen por fase
    rs = lib.wb.create_sheet("Resumen por fase")
    estilo_titulo(rs, "Resumen de paquetes por fase del proyecto", f"{dp.PROYECTO}. Calculado a partir de la hoja «Paquetes».", 8)
    R = {k: f"Paquetes!{rango(c, k, ult)}" for k in ("tipo", "fase", "estado", "situacion", "avance", "restr", "hh")}
    fila = 4
    for fase in dp.FASES:
        def cnt(extra, fase=fase):
            return lambda ref: f'=COUNTIFS({R["tipo"]},{ref},{R["fase"]},"{fase}"{extra})'
        columnas = [
            ("Total", cnt(""), "0"),
            ("Planificados / en desarrollo", lambda ref, fase=fase: f'=COUNTIFS({R["tipo"]},{ref},{R["fase"]},"{fase}",{R["estado"]},"Planificado")+COUNTIFS({R["tipo"]},{ref},{R["fase"]},"{fase}",{R["estado"]},"En desarrollo")+COUNTIFS({R["tipo"]},{ref},{R["fase"]},"{fase}",{R["estado"]},"En revisión")', "0"),
            ("Aprobados / liberados", lambda ref, fase=fase: f'=COUNTIFS({R["tipo"]},{ref},{R["fase"]},"{fase}",{R["estado"]},"Aprobado")+COUNTIFS({R["tipo"]},{ref},{R["fase"]},"{fase}",{R["estado"]},"Liberado")', "0"),
            ("En ejecución", cnt(f',{R["estado"]},"En ejecución"'), "0"),
            ("Cerrados", cnt(f',{R["estado"]},"Cerrado"'), "0"),
            ("Atrasados", lambda ref, fase=fase: f'=COUNTIFS({R["tipo"]},{ref},{R["fase"]},"{fase}",{R["situacion"]},"Atrasado")+COUNTIFS({R["tipo"]},{ref},{R["fase"]},"{fase}",{R["situacion"]},"Inicio atrasado")', "0"),
            ("Restricciones abiertas", lambda ref, fase=fase: f'=SUMIFS({R["restr"]},{R["tipo"]},{ref},{R["fase"]},"{fase}")', "0"),
            ("% avance promedio", lambda ref, fase=fase: f'=IFERROR(AVERAGEIFS({R["avance"]},{R["tipo"]},{ref},{R["fase"]},"{fase}"),"")', "0%"),
        ]
        _, fila = tabla_resumen(rs, fila, 1, f"{fase} – {dp.NOMBRE_FASE[fase]}", dp.TIPOS_PAQUETE, columnas, 14)
        fila += 1
    for i in range(2, 10):
        rs.column_dimensions[get_column_letter(i)].width = 14
    configurar_impresion(rs, "Resumen por fase", area=f"A1:I{fila}")
    lib.wb.move_sheet("Resumen por fase", offset=-1)

    estados = "\n".join(f"• {k}: {v}" for k, v in dp.SIGNIFICADO_ESTADO_PAQUETE.items())
    hoja_instrucciones(
        lib,
        "Registrar todos los paquetes AWP del proyecto (CWA, CWP, EWP, PWP e IWP) con su fase, sus relaciones, estado, fechas plan y real, avance y restricciones abiertas, para controlar que ingeniería y procura entreguen a tiempo lo que construcción necesita. Es el equivalente del «Pack Track» del procedimiento 3.0.",
        "Desde la planificación temprana (cuando se define el plan de liberación CWP/EWP/PWP) hasta el cierre de cada fase. Se actualiza cada semana.",
        "Controles del proyecto y el Líder de WFP. Ingeniería actualiza sus EWP, procura sus PWP y los planificadores sus IWP.",
        [
            "Actualice la fecha de corte (celda B3). Para que se actualice sola, escriba =HOY().",
            "Borre las filas amarillas de ejemplo.",
            "Registre primero las CWA, luego los CWP y, para cada CWP, su EWP y su PWP (mismo sufijo de código). Agregue los IWP a medida que se crean.",
            "Complete siempre las columnas CWA y CWP asociado: son las que relacionan los paquetes entre sí.",
            "Cargue fechas plan desde el Path of Construction y el programa de liberación de IWP; no cambie la línea base, registre las fechas reales.",
            "Pegue en la hoja «Restricciones» las columnas del registro de restricciones para calcular las restricciones abiertas.",
            "Revise la hoja «Resumen por fase» para comparar las fases.",
        ],
        [("Paquetes", cols), ("Restricciones", rcols)],
        notas=[("Significado de los estados", estados),
               ("Formato condicional", "Situación en rojo: paquete atrasado o con inicio atrasado. Ámbar: cerrado con atraso o con restricciones abiertas. Verde: en plazo o cerrado sin atraso. La barra azul muestra el % de avance.")],
    )
    lib.crear_listas()
    lib.guardar()


# ==========================================================================
# 4. Definición de CWA
# ==========================================================================

def definicion_cwa():
    lib = Libro(SALIDA / "Definicion_CWA.xlsx", "Definición de CWA", ["Fases", "EstadosPaquete", "Roles", "SiNo"])
    fechas = {
        "CWA-1.01": (D(2027, 6, 7), D(2027, 12, 20), "Cerrado", 4), "CWA-1.02": (D(2027, 8, 2), D(2028, 1, 31), "Cerrado", 5),
        "CWA-1.03": (D(2027, 5, 3), D(2027, 11, 29), "Cerrado", 3), "CWA-2.01": (D(2028, 2, 7), D(2029, 2, 26), "En ejecución", 14),
        "CWA-2.02": (D(2028, 5, 1), D(2029, 3, 26), "Aprobado", 13), "CWA-2.03": (D(2028, 1, 10), D(2028, 12, 18), "En ejecución", 8),
        "CWA-2.04": (D(2028, 7, 3), D(2029, 4, 30), "Aprobado", 10), "CWA-3.01": (D(2029, 2, 5), D(2029, 10, 29), "En desarrollo", 8),
        "CWA-3.02": (D(2029, 2, 5), D(2029, 10, 29), "En desarrollo", 9), "CWA-3.03": (D(2029, 6, 4), D(2029, 11, 26), "En desarrollo", 3),
    }
    filas = []
    for cod, fase, nombre, disc, hh, sec, pred, lim in dp.CWAS:
        ini, fin, est, ncwp = fechas[cod]
        filas.append(dict(ejemplo=EJ, cod=cod, fase=fase, nombre=nombre, lim=lim, disc=disc, hh=hh, sec=sec, pred=pred,
                          ncwp=ncwp, ini=ini, fin=fin, rol="Gerente de Construcción", estado=est, crit="Sí", obs=""))

    def pct(f, c, p):
        return f'=IF({c("cod")}{f}="","",IFERROR({c("hh")}{f}/SUMIFS({rango(c, "hh", ult_)},{rango(c, "fase", ult_)},{c("fase")}{f}),""))'

    def dur(f, c, p):
        return f'=IF(OR({c("ini")}{f}="",{c("fin")}{f}=""),"",ROUND(({c("fin")}{f}-{c("ini")}{f})/30.4,1))'

    def hhcwp(f, c, p):
        return f'=IF(OR({c("hh")}{f}="",{c("ncwp")}{f}="",{c("ncwp")}{f}=0),"",{c("hh")}{f}/{c("ncwp")}{f})'

    def valida(f, c, p):
        return f'=IF({c("cod")}{f}="","",IF({c("hhcwp")}{f}="","",IF({c("hhcwp")}{f}>40000,"Revisar: CWP > 40 000 HH","OK")))'

    ult_ = FILA_DATOS + 100 - 1
    cols = [
        COL_EJEMPLO,
        Col("cod", "Código CWA", 10, desc="Código CWA-<fase>.<nn>. Una CWA pertenece a una sola fase."),
        Col("fase", "Fase del proyecto", 11, lista="Fases", desc="Fase del proyecto."),
        Col("nombre", "Nombre", 28, ajustar=True, desc="Nombre corto del área."),
        Col("lim", "Límites (ejes, coordenadas, referencias)", 34, ajustar=True, desc="Descripción de los límites físicos o lógicos. Deben coincidir con el plano de CWA."),
        Col("disc", "Disciplinas", 14, desc="Códigos de las disciplinas que trabajan en la CWA."),
        Col("hh", "HH estimadas", 11, "num", desc="Horas-hombre directas estimadas de construcción."),
        Col("pct", "% HH de la fase", 9, "pct", formula=pct, desc="Calculado: participación de la CWA en las HH de su fase."),
        Col("sec", "Secuencia en el PoC", 9, "num", desc="Orden de la CWA dentro del Path of Construction de su fase."),
        Col("pred", "CWA predecesora", 11, desc="CWA que debe terminar (o entregar un área) antes de iniciar esta."),
        Col("ncwp", "N.° de CWP previstos", 9, "num", desc="Cantidad estimada de CWP (uno por disciplina y frente)."),
        Col("hhcwp", "HH promedio por CWP", 11, "num", formula=hhcwp, desc="Calculado: HH de la CWA ÷ N.° de CWP."),
        Col("valida", "Control tamaño CWP", 17, formula=valida, desc="Calculado: avisa si el CWP promedio supera 40 000 HH (criterio de las fuentes)."),
        Col("ini", "Inicio plan", 11, "fecha", desc="Inicio de construcción previsto en la CWA."),
        Col("fin", "Fin plan", 11, "fecha", desc="Fin de construcción previsto."),
        Col("dur", "Duración (meses)", 9, "dec", formula=dur, desc="Calculado: duración aproximada en meses."),
        Col("rol", "Responsable (rol)", 20, lista="Roles", desc="Rol responsable de la CWA en construcción."),
        Col("estado", "Estado", 12, lista="EstadosPaquete", desc="Estado de la CWA."),
        Col("crit", "Cumple criterios", 9, lista="SiNo", desc="«Sí» si la CWA cumple la lista de la hoja «Criterios»."),
        Col("obs", "Observaciones", 28, ajustar=True, desc="Comentarios."),
    ]
    ws, c, ult = hoja_datos(lib, "CWA", "Definición de áreas de trabajo de construcción (CWA)", cols, filas, n_filas=100)
    v = f"${c('valida')}{FILA_DATOS}"
    resaltar_fila(ws, f"{c('valida')}{FILA_DATOS}:{c('valida')}{ult}", f'LEFT({v},7)="Revisar"', ROJO, ROJO_TEXTO)
    resaltar_fila(ws, f"{c('valida')}{FILA_DATOS}:{c('valida')}{ult}", f'{v}="OK"', VERDE, VERDE_TEXTO)
    cr = f"${c('crit')}{FILA_DATOS}"
    resaltar_fila(ws, f"{c('crit')}{FILA_DATOS}:{c('crit')}{ult}", f'{cr}="No"', ROJO, ROJO_TEXTO)

    # Criterios
    cs = lib.wb.create_sheet("Criterios")
    estilo_titulo(cs, "Criterios para definir una CWA", "Revise cada CWA con esta lista antes del hito H1 (CWA definidas).", 3)
    criterios = [
        ("Límites", "Tiene límites físicos o lógicos claros, dibujados en un plano de CWA y sin superposición con otras CWA."),
        ("Fase", "Pertenece a una sola fase del proyecto."),
        ("Secuencia", "Permite construir en una secuencia lógica que termina en sistemas o áreas entregables."),
        ("Tamaño", "Su tamaño permite dividirla en CWP de una disciplina y menos de 40 000 HH."),
        ("Accesos", "Considera accesos, zonas de acopio, grúas y rutas de izaje."),
        ("Interfaces", "Identifica las interfaces con CWA vecinas y con otras fases (entregas de área, conexiones)."),
        ("Sistemas", "Identifica los sistemas de comisionamiento que atraviesan la CWA."),
        ("Modularización", "Evalúa si parte del alcance se prefabrica o modulariza."),
        ("Código", "Tiene código único según la codificación del kit y está cargada en el modelo 3D y el cronograma."),
        ("Aprobación", "Fue revisada por construcción, ingeniería, procura y comisionamiento y aprobada por el AWP Champion."),
    ]
    for i, t in enumerate(("N.°", "Criterio", "Descripción"), 1):
        celda = cs.cell(row=4, column=i, value=t)
        celda.font = Font(name=FUENTE, bold=True, color=BLANCO)
        celda.fill = relleno(AZUL)
        celda.border = BORDE
    for n, (k, d) in enumerate(criterios, 1):
        for i, v in enumerate((n, k, d), 1):
            celda = cs.cell(row=4 + n, column=i, value=v)
            celda.border = BORDE
            celda.alignment = Alignment(wrap_text=True, vertical="top")
    cs.column_dimensions["A"].width = 6
    cs.column_dimensions["B"].width = 18
    cs.column_dimensions["C"].width = 90
    configurar_impresion(cs, "Criterios de CWA", "landscape", "4:4", f"A1:C{4 + len(criterios)}")

    hoja_instrucciones(
        lib,
        "Definir y documentar las áreas de trabajo de construcción (CWA) de cada fase: límites, disciplinas, tamaño, secuencia y predecesoras. Las CWA son la base de todo el flujo de paquetes y del Path of Construction.",
        "En la planificación temprana de cada fase (hito H1). Se revisa si cambia el alcance o la secuencia.",
        "El AWP Champion con el Gerente de Construcción; participan ingeniería, procura y comisionamiento.",
        [
            "Borre las filas amarillas de ejemplo.",
            "Registre una fila por CWA, con su código y fase.",
            "Describa los límites de forma que coincidan con el plano de CWA.",
            "Estime las HH y la cantidad de CWP; revise el control de tamaño.",
            "Indique la secuencia y la CWA predecesora según el Path of Construction.",
            "Revise cada CWA con la hoja «Criterios» y marque «Cumple criterios».",
        ],
        [("CWA", cols)],
        justificacion="Las fuentes (procedimiento 1.0, curso 2023 y marco educativo del CII) sitúan la definición de CWA como el primer entregable AWP de la planificación temprana. Sin una lista formal de CWA no se pueden codificar ni secuenciar los paquetes.",
    )
    lib.crear_listas()
    lib.guardar()


# ==========================================================================
# 5. Path of Construction
# ==========================================================================

def path_of_construction():
    lib = Libro(SALIDA / "Path_of_Construction.xlsx", "Path of Construction", ["Fases", "Disciplinas", "SiNo"])
    filas = []
    for n, (cod, fase, desc, disc, hh, ini, dur, pred, sistema, just) in enumerate(dp.CWPS, 1):
        filas.append(dict(ejemplo=EJ, sec=n, fase=fase, cwp=cod, desc=desc, disc=disc, hh=hh, pred=pred, ini=ini, dur=dur, sistema=sistema, just=just))
    ult_ = FILA_DATOS + 150 - 1

    def cwa(f, c, p):
        return f'=IF({c("cwp")}{f}="","","CWA-"&MID({c("cwp")}{f},5,4))'

    def fin(f, c, p):
        return f'=IF(OR({c("ini")}{f}="",{c("dur")}{f}=""),"",{c("ini")}{f}+7*{c("dur")}{f}-1)'

    def menos(par, factor):
        def f(fi, c, p):
            return f'=IF({c("ini")}{fi}="","",{c("ini")}{fi}-{factor}*{p[par]})'
        return f

    def control(f, c, p):
        pr = f'{c("pred")}{f}'
        return (f'=IF(OR({c("cwp")}{f}="",{pr}="",{pr}="—"),IF({c("cwp")}{f}="","","OK"),'
                f'IFERROR(IF(INDEX({rango(c, "fin", ult_)},MATCH({pr},{rango(c, "cwp", ult_)},0))>={c("ini")}{f},"Conflicto con predecesor","OK"),"Predecesor no encontrado"))')

    cols = [
        COL_EJEMPLO,
        Col("sec", "Secuencia", 8, "num", desc="Orden del CWP en el Path of Construction."),
        Col("fase", "Fase del proyecto", 11, lista="Fases", desc="Fase del proyecto."),
        Col("cwa", "CWA", 10, formula=cwa, desc="Calculado a partir del código del CWP."),
        Col("cwp", "CWP", 17, desc="Código del CWP."),
        Col("desc", "Descripción", 30, ajustar=True, desc="Alcance resumido del CWP."),
        Col("disc", "Disciplina", 9, lista="Disciplinas", desc="Disciplina del CWP."),
        Col("hh", "HH estimadas", 10, "num", desc="Horas-hombre directas estimadas."),
        Col("pred", "Predecesor", 17, desc="CWP que debe terminar antes (relación fin–comienzo). «—» si no tiene."),
        Col("ini", "Inicio plan", 11, "fecha", desc="Fecha de inicio de construcción del CWP."),
        Col("dur", "Duración (semanas)", 9, "num", desc="Duración de construcción en semanas."),
        Col("fin", "Fin plan", 11, "fecha", formula=fin, desc="Calculado: inicio + duración."),
        Col("f_ewp", "EWP IFC requerido", 11, "fecha", formula=menos("ewp", 7), desc="Calculado: fecha límite para emitir el EWP IFC (inicio − anticipación de EWP)."),
        Col("f_cwp", "Emisión del CWP", 11, "fecha", formula=menos("cwp", 7), desc="Calculado: fecha de emisión del CWP (inicio − anticipación de CWP)."),
        Col("f_ras", "Materiales en obra (RAS)", 11, "fecha", formula=menos("ras", 1), desc="Calculado: fecha requerida en obra de los materiales del PWP (inicio − días RAS)."),
        Col("control", "Control de secuencia", 18, formula=control, desc="Calculado: «Conflicto» si el predecesor termina después del inicio de este CWP."),
        Col("sistema", "Sistema", 16, desc="Sistema de comisionamiento que completa el CWP."),
        Col("just", "Justificación de la secuencia", 40, ajustar=True, desc="Por qué el CWP va en esa posición (acuerdo del taller IPP)."),
    ]
    params = [("m1", "Mes M1 del proyecto:", dp.INICIO_PROYECTO, "fecha"),
              ("ewp", "Anticipación EWP (semanas):", 13, None),
              ("cwp", "Anticipación CWP (semanas):", 10, None),
              ("ras", "Anticipación RAS (días):", 10, None)]
    # 36 columnas de meses para el diagrama de barras
    base = len(cols)
    for k in range(1, 37):
        def mes(f, c, p, k=k):
            col = get_column_letter(base + k)
            return (f'=IF(OR({c("ini")}{f}="",{c("fin")}{f}=""),"",IF(AND({c("ini")}{f}<=EOMONTH({col}$4,0),'
                    f'{c("fin")}{f}>={col}$4),1,""))')
        cols.append(Col(f"m{k}", f"M{k}", 3.2, "num", formula=mes, desc=""))
    ws, c, ult = hoja_datos(lib, "Path of Construction", "Path of Construction (secuencia de CWP)", cols, filas, n_filas=150,
                            params=params, congelar_col=5)
    for k in range(1, 37):
        col = c(f"m{k}")
        ws[f"{col}4"] = f"=EDATE({ws._kit_celdas_param['m1']},{k - 1})"
        ws[f"{col}4"].number_format = "mm/yy"
        ws[f"{col}4"].font = Font(name=FUENTE, size=7, color="595959")
        ws[f"{col}4"].alignment = Alignment(text_rotation=90, horizontal="center")
        ws[f"{col}{FILA_ENC}"].alignment = Alignment(text_rotation=90, horizontal="center", vertical="center")
        for fila in range(FILA_DATOS, ult + 1):
            ws[f"{col}{fila}"].fill = relleno(BLANCO)
    ws.row_dimensions[4].height = 34
    rm = f"{c('m1')}{FILA_DATOS}:{c('m36')}{ult}"
    resaltar_fila(ws, rm, f"{c('m1')}{FILA_DATOS}=1", "2F5597", "2F5597")
    ctl = f"${c('control')}{FILA_DATOS}"
    resaltar_fila(ws, f"{c('control')}{FILA_DATOS}:{c('control')}{ult}", f'AND({ctl}<>"",{ctl}<>"OK")', ROJO, ROJO_TEXTO)
    resaltar_fila(ws, f"{c('control')}{FILA_DATOS}:{c('control')}{ult}", f'{ctl}="OK"', VERDE, VERDE_TEXTO)
    ws.auto_filter.ref = f"A{FILA_ENC}:{c('just')}{ult}"

    hoja_instrucciones(
        lib,
        "Documentar el Path of Construction (PoC): la secuencia en que se construirán los CWP de cada fase, con sus predecesores, y calcular hacia atrás las fechas que ingeniería (EWP), construcción (emisión de CWP) y procura (RAS) deben cumplir.",
        "Se desarrolla en los talleres de planificación interactiva (IPP) de la planificación temprana de cada fase y se congela en el hito H2. Se revisa ante cambios mayores.",
        "El Gerente de Construcción lidera; el AWP Champion facilita; ingeniería, procura, comisionamiento y controles del proyecto participan.",
        [
            "Ajuste los parámetros de la fila 3: mes M1 del proyecto y anticipaciones de EWP, CWP y RAS.",
            "Borre las filas amarillas de ejemplo.",
            "Registre los CWP en el orden acordado en el taller, con su predecesor, inicio y duración.",
            "Revise la columna «Control de secuencia»: no debe haber conflictos.",
            "Traslade las fechas «EWP IFC requerido» y «RAS» al plan de ingeniería y de procura, y a Seguimiento_Paquetes.",
            "Las columnas M1 a M36 muestran un diagrama de barras por mes (no las edite).",
        ],
        [("Path of Construction", [col for col in cols if not col.clave.startswith("m")] +
          [("M1 … M36", "Calculado: meses del proyecto en que se construye el CWP (barra azul).", "Calculado")])],
        justificacion="Las fuentes coinciden en que la secuencia de construcción debe dirigir la ingeniería y la procura. Esta plantilla convierte el PoC en fechas requeridas concretas y detecta conflictos de secuencia.",
    )
    lib.crear_listas()
    lib.guardar()


# ==========================================================================
# 6. Programa de liberación de IWP
# ==========================================================================

PLAZOS = [  # fase, IWP iniciado, restricciones identificadas, asignadas, levantadas, liberación (semanas antes)
    ("Fase 1", 8, 6, 5, 3, 2), ("Fase 2", 12, 10, 8, 4, 2), ("Fase 3", 8, 6, 5, 3, 2),
]


def programa_liberacion():
    lib = Libro(SALIDA / "Programa_Liberacion_IWP.xlsx", "Programa de liberación de IWP", ["Fases", "EstadosPaquete", "Roles"])
    iwps = [
        ("IWP-1.01-CIV-01-001", "Fase 1", "Excavación masiva ejes 1–5", "Cuadrilla C-01", 480, D(2027, 6, 7), D(2027, 6, 11), 0, "Cerrado"),
        ("IWP-1.02-TUB-01-001", "Fase 1", "Tubería de agua tramo T1", "Cuadrilla T-02", 360, D(2027, 8, 2), D(2027, 7, 19), 0, "Cerrado"),
        ("IWP-1.02-ELE-01-001", "Fase 1", "Banco de ductos BD-01 a BD-04", "Cuadrilla E-01", 400, D(2027, 8, 16), D(2027, 8, 6), 0, "Cerrado"),
        ("IWP-2.01-CIV-01-001", "Fase 2", "Zapatas ejes A1–A4", "Cuadrilla C-05", 520, D(2028, 2, 7), D(2028, 1, 26), 0, "Cerrado"),
        ("IWP-2.01-EST-01-001", "Fase 2", "Montaje de columnas nivel 1 ejes A1–A4", "Cuadrilla M-01", 420, D(2028, 4, 25), None, 2, "En desarrollo"),
        ("IWP-2.03-ELE-01-001", "Fase 2", "Montaje de tableros de media tensión", "Cuadrilla E-04", 380, D(2028, 7, 3), None, 1, "En desarrollo"),
        ("IWP-2.01-TUB-01-001", "Fase 2", "Tubería de agua de servicio nivel 1", "Cuadrilla T-06", 450, D(2028, 8, 28), None, 1, "En desarrollo"),
        ("IWP-3.02-CIV-01-001", "Fase 3", "Zapatas ejes C1–C3", "Por asignar", 500, D(2029, 2, 5), None, 0, "Planificado"),
        ("IWP-3.01-ARQ-01-001", "Fase 3", "Tabiquería nivel 1", "Por asignar", 350, D(2029, 3, 5), None, 1, "Planificado"),
    ]
    filas = [dict(ejemplo=EJ, iwp=i, fase=f, desc=d, cuad=q, hh=h, ini=s, lib=l, restr=r, estado=e) for i, f, d, q, h, s, l, r, e in iwps]
    tabla = "Plazos!$A$6:$F$8"

    def hito(n):
        def f(fi, c, p):
            return f'=IF(OR({c("iwp")}{fi}="",{c("ini")}{fi}="",{c("fase")}{fi}=""),"",{c("ini")}{fi}-7*VLOOKUP({c("fase")}{fi},{tabla},{n},FALSE))'
        return f

    def cwp(f, c, p):
        return f'=IF({c("iwp")}{f}="","","CWP"&MID({c("iwp")}{f},4,12))'

    def atraso(f, c, p):
        o, l = f'{c("f_lib")}{f}', f'{c("lib")}{f}'
        return f'=IF(OR({c("iwp")}{f}="",{o}=""),"",IF({l}<>"",MAX(0,{l}-{o}),MAX(0,{p["corte"]}-{o})))'

    def sit(f, c, p):
        o, l, a = f'{c("f_lib")}{f}', f'{c("lib")}{f}', f'{c("atraso")}{f}'
        return (f'=IF(OR({c("iwp")}{f}="",{o}=""),"",IF({l}<>"",IF({a}>0,"Liberado con atraso","Liberado a tiempo"),'
                f'IF({p["corte"]}>{o},"Vencido sin liberar",IF({o}-{p["corte"]}<=14,"Próximo a liberar","En plazo"))))')

    cols = [
        COL_EJEMPLO,
        Col("iwp", "IWP", 21, desc="Código del IWP."),
        Col("fase", "Fase del proyecto", 11, lista="Fases", desc="Fase del proyecto (define los plazos de la hoja «Plazos»)."),
        Col("cwp", "CWP", 17, formula=cwp, desc="Calculado a partir del código del IWP."),
        Col("desc", "Descripción", 30, ajustar=True, desc="Alcance del IWP."),
        Col("cuad", "Cuadrilla / capataz", 14, desc="Cuadrilla o capataz asignado."),
        Col("hh", "HH estimadas", 9, "num", desc="Horas-hombre del IWP (300 a 600 HH)."),
        Col("ini", "Inicio de ejecución plan", 12, "fecha", desc="Fecha planificada de inicio en campo (cronograma de nivel 5)."),
        Col("f_ini", "IWP iniciado", 11, "fecha", formula=hito(2), desc="Calculado: fecha en que debe definirse el alcance del IWP."),
        Col("f_ident", "Restricciones identificadas", 11, "fecha", formula=hito(3), desc="Calculado: fecha límite para identificar restricciones."),
        Col("f_asig", "Restricciones asignadas", 11, "fecha", formula=hito(4), desc="Calculado: fecha límite para asignar responsables."),
        Col("f_lev", "Restricciones levantadas", 11, "fecha", formula=hito(5), desc="Calculado: fecha límite para levantar todas las restricciones."),
        Col("f_lib", "Liberación objetivo", 11, "fecha", formula=hito(6), desc="Calculado: fecha objetivo de liberación (ingreso al backlog)."),
        Col("lib", "Liberación real", 11, "fecha", desc="Fecha real de liberación con checklist firmado."),
        Col("restr", "Restricciones abiertas", 10, "num", desc="Restricciones abiertas o en gestión (del registro de restricciones)."),
        Col("estado", "Estado", 12, lista="EstadosPaquete", desc="Estado del IWP."),
        Col("atraso", "Días de atraso en liberación", 10, "num", formula=atraso, desc="Calculado: días de atraso respecto de la liberación objetivo."),
        Col("sit", "Situación", 18, formula=sit, desc="Calculado: Liberado a tiempo, Liberado con atraso, Vencido sin liberar, Próximo a liberar (≤ 14 días) o En plazo."),
        Col("obs", "Observaciones", 26, ajustar=True, desc="Comentarios."),
    ]
    ws, c, ult = hoja_datos(lib, "Programa", "Programa de liberación de IWP", cols, filas, n_filas=300,
                            params=[("corte", "Fecha de corte:", dp.FECHA_CORTE, "fecha")])
    s = f"${c('sit')}{FILA_DATOS}"
    todo = f"A{FILA_DATOS}:{c('obs')}{ult}"
    resaltar_fila(ws, todo, f'{s}="Vencido sin liberar"', ROJO, ROJO_TEXTO)
    resaltar_fila(ws, todo, f'{s}="Próximo a liberar"', AMBAR)
    resaltar_fila(ws, f"{c('sit')}{FILA_DATOS}:{c('sit')}{ult}", f'LEFT({s},8)="Liberado"', VERDE, VERDE_TEXTO)

    ps = lib.wb.create_sheet("Plazos")
    estilo_titulo(ps, "Plazos de liberación por fase (semanas antes del inicio de ejecución)",
                  "Basado en el programa típico de restricciones del marco educativo del CII. Puede ajustarlos: las fechas del programa se recalculan.", 6)
    enc = ["Fase del proyecto", "IWP iniciado", "Restricciones identificadas", "Restricciones asignadas", "Restricciones levantadas", "Liberación (backlog)"]
    for i, t in enumerate(enc, 1):
        celda = ps.cell(row=5, column=i, value=t)
        celda.font = Font(name=FUENTE, bold=True, color=BLANCO)
        celda.fill = relleno(AZUL)
        celda.border = BORDE
        celda.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        ps.column_dimensions[get_column_letter(i)].width = 16
    ps.row_dimensions[5].height = 32
    for n, fila in enumerate(PLAZOS, 6):
        for i, v in enumerate(fila, 1):
            celda = ps.cell(row=n, column=i, value=v)
            celda.border = BORDE
            celda.alignment = Alignment(horizontal="center")
            if i > 1:
                celda.fill = relleno(AZUL_CLARO)
    configurar_impresion(ps, "Plazos por fase", "landscape", area="A1:F8")

    hoja_instrucciones(
        lib,
        "Programar hacia atrás, para cada IWP, las fechas en que debe iniciarse, en que deben identificarse, asignarse y levantarse sus restricciones, y en que debe liberarse, según los plazos de su fase. Permite ver con anticipación qué IWP no llegarán liberados a tiempo.",
        "Cada semana, para los IWP que entran en la ventana de planificación (8 semanas en las fases 1 y 3, 12 semanas en la fase 2).",
        "Los planificadores de frente de trabajo; el Líder de WFP revisa y usa la situación en la reunión de restricciones.",
        [
            "Actualice la fecha de corte (B3).",
            "Revise los plazos por fase en la hoja «Plazos».",
            "Borre las filas amarillas de ejemplo y registre los IWP con su fecha de inicio de ejecución planificada.",
            "Las fechas intermedias se calculan solas; registre la liberación real cuando el checklist esté firmado.",
            "Copie las restricciones abiertas del registro de restricciones (o de Seguimiento_Paquetes).",
            "Filtre por «Situación» para preparar la reunión semanal.",
        ],
        [("Programa", cols)],
        justificacion="El procedimiento 3.0 (Pack Track) y el marco educativo del CII definen hitos de preparación del IWP en semanas antes de la ejecución. Esta plantilla aplica esos plazos a cada fase del proyecto tipo.",
        notas=[("Formato condicional", "Rojo: IWP con liberación objetivo vencida y sin liberar. Ámbar: liberación objetivo en 14 días o menos. Verde: liberado.")],
    )
    lib.crear_listas()
    lib.guardar()


# ==========================================================================
# 7. Lookahead de 3 semanas
# ==========================================================================

def lookahead():
    causas = dp.CAUSAS_NO_CUMPLIMIENTO
    lib = Libro(SALIDA / "Lookahead_3_Semanas.xlsx", "Lookahead de 3 semanas", ["Fases", "SiNo"])
    inicio = D(2029, 3, 5)
    datos = [
        ("Fase 2", "IWP-2.01-TUB-01-021", "Tubería de agua de servicio nivel 3 ejes A5–A8", "Superint. tuberías / Capataz T-06", 8, 440, 0, "Andamio AN-31", "111111000000000000", "Sí", ""),
        ("Fase 2", "IWP-2.01-TUB-01-022", "Tubería de agua de servicio nivel 3 ejes A9–A12", "Superint. tuberías / Capataz T-06", 8, 440, 0, "Andamio AN-32", "000000111111000000", "", ""),
        ("Fase 2", "IWP-2.03-ELE-01-012", "Tendido de cables de media tensión al bloque A", "Superint. eléctrico / Capataz E-04", 10, 560, 0, "Carretes y tiracables", "111110000000000000", "No", "Materiales"),
        ("Fase 2", "IWP-2.03-ELE-01-013", "Conexionado de tableros de media tensión", "Superint. eléctrico / Capataz E-04", 6, 300, 0, "Equipo de pruebas", "000000111111000000", "", ""),
        ("Fase 3", "IWP-3.02-CIV-01-003", "Zapatas ejes C4–C6", "Superint. civil / Capataz C-08", 12, 580, 0, "Bomba de concreto", "111111000000000000", "Sí", ""),
        ("Fase 3", "IWP-3.02-CIV-01-004", "Pedestales ejes C1–C3", "Superint. civil / Capataz C-08", 10, 420, 0, "Encofrado metálico", "000000111111000000", "", ""),
        ("Fase 3", "IWP-3.01-ARQ-01-001", "Tabiquería nivel 1 edificio de oficinas", "Superint. arquitectura / Capataz A-02", 9, 350, 1, "Plataforma elevadora", "000000000000111111", "", ""),
        ("Fase 3", "IWP-3.01-ARQ-01-002", "Tabiquería nivel 2 edificio de oficinas", "Superint. arquitectura / Capataz A-02", 9, 350, 0, "Plataforma elevadora", "", "", ""),
    ]
    dias = [f"d{k}" for k in range(18)]
    filas = []
    for fase, iwp, desc, resp, cuad, hh, restr, rec, prog, comp, causa in datos:
        f = dict(ejemplo=EJ, fase=fase, iwp=iwp, desc=desc, resp=resp, cuad=cuad, hh=hh, restr=restr, rec=rec, comp=comp, causa=causa)
        for k, ch in enumerate(prog):
            if ch == "1":
                f[dias[k]] = "X"
        filas.append(f)

    def cwp(f, c, p):
        return f'=IF({c("iwp")}{f}="","","CWP"&MID({c("iwp")}{f},4,12))'

    def listo(f, c, p):
        return f'=IF({c("iwp")}{f}="","",IF({c("restr")}{f}=0,"Sí","No"))'

    def plan_s1(f, c, p):
        return f'=IF({c("iwp")}{f}="","",IF(COUNTIF({c("d0")}{f}:{c("d5")}{f},"X")>0,"Sí","No"))'

    cols = [
        COL_EJEMPLO,
        Col("fase", "Fase del proyecto", 10, lista="Fases", desc="Fase del proyecto."),
        Col("cwp", "CWP", 16, formula=cwp, desc="Calculado a partir del código del IWP."),
        Col("iwp", "IWP", 20, desc="Código del IWP seleccionado del backlog."),
        Col("desc", "Descripción", 30, ajustar=True, desc="Alcance del IWP."),
        Col("resp", "Superintendente / capataz", 22, ajustar=True, desc="Responsables de la ejecución."),
        Col("cuad", "Personas en cuadrilla", 8, "num", desc="Tamaño de la cuadrilla."),
        Col("hh", "HH", 7, "num", desc="Horas-hombre del IWP."),
        Col("restr", "Restricciones abiertas", 9, "num", desc="Restricciones abiertas del IWP (del registro)."),
        Col("listo", "Listo para ejecutar", 8, formula=listo, desc="Calculado: «Sí» si no tiene restricciones abiertas."),
        Col("rec", "Recursos críticos", 18, ajustar=True, desc="Grúas, andamios, equipos o permisos que deben confirmarse."),
    ]
    etiquetas = ["L", "M", "X", "J", "V", "S"]
    for k in range(18):
        cols.append(Col(dias[k], f"S{k // 6 + 1} {etiquetas[k % 6]}", 3.6, lista_valores=["X"], desc=""))
    cols += [
        Col("plan", "Planificado semana 1", 9, formula=plan_s1, desc="Calculado: «Sí» si el IWP tiene días marcados en la semana 1."),
        Col("comp", "Completado semana 1", 9, lista="SiNo", desc="Al cierre de la semana 1: «Sí» si el IWP planificado se completó."),
        Col("causa", "Causa de no cumplimiento", 20, lista="Causas", desc="Si no se completó, causa principal (sirve para el análisis de PPC)."),
    ]
    ws, c, ult = hoja_datos(
        lib, "Lookahead", "Lookahead de 3 semanas", cols, filas, n_filas=120,
        params=[("inicio", "Lunes de la semana 1:", inicio, "fecha"), ("ppc", "PPC semana 1:", None, "0%")],
        subtitulo=f"{dp.PROYECTO}. Ejemplo en M27, con las fases 2 y 3 en construcción simultánea (la fase 1 ya está cerrada). Filas amarillas = ejemplo.",
        congelar_col=4,
    )
    p = ws._kit_celdas_param
    ws[p["ppc"].replace("$", "")] = (f'=IFERROR(COUNTIFS({rango(c, "plan", ult)},"Sí",{rango(c, "comp", ult)},"Sí")/'
                                     f'COUNTIF({rango(c, "plan", ult)},"Sí"),"")')
    for k in range(18):
        col = c(dias[k])
        ws[f"{col}4"] = f"={p['inicio']}+{(k // 6) * 7 + k % 6}"
        ws[f"{col}4"].number_format = "dd/mm"
        ws[f"{col}4"].font = Font(name=FUENTE, size=7, color="595959")
        ws[f"{col}4"].alignment = Alignment(text_rotation=90, horizontal="center")
        ws[f"{col}{FILA_ENC}"].alignment = Alignment(text_rotation=90, horizontal="center", vertical="center")
        for fila in range(FILA_DATOS, ult + 1):
            ws[f"{col}{fila}"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[4].height = 30
    rd = f"{c('d0')}{FILA_DATOS}:{c('d17')}{ult}"
    celda = f"{c('d0')}{FILA_DATOS}"
    l = f"${c('listo')}{FILA_DATOS}"
    resaltar_fila(ws, rd, f'AND({celda}="X",{l}="No")', "C00000", BLANCO, True)
    resaltar_fila(ws, rd, f'{celda}="X"', "2F5597", BLANCO, True)
    resaltar_fila(ws, f"{c('listo')}{FILA_DATOS}:{c('listo')}{ult}", f'{l}="No"', ROJO, ROJO_TEXTO)
    resaltar_fila(ws, f"{c('listo')}{FILA_DATOS}:{c('listo')}{ult}", f'{l}="Sí"', VERDE, VERDE_TEXTO)
    cp = f"${c('comp')}{FILA_DATOS}"
    resaltar_fila(ws, f"{c('comp')}{FILA_DATOS}:{c('causa')}{ult}", f'AND(${c("plan")}{FILA_DATOS}="Sí",{cp}="No")', AMBAR)

    # Resumen de causas
    rs = lib.wb.create_sheet("Causas de no cumplimiento")
    estilo_titulo(rs, "Causas de no cumplimiento (semana 1)", "Calculado a partir de la hoja «Lookahead». Úselo en la reunión semanal y en el análisis de lecciones.", 3)
    rf, rc = f"Lookahead!{rango(c, 'fase', ult)}", f"Lookahead!{rango(c, 'causa', ult)}"
    tabla_resumen(rs, 4, 1, "IWP no completados por causa y fase", causas,
                  [(f, (lambda ref, f=f: f'=COUNTIFS({rc},{ref},{rf},"{f}")'), "0") for f in dp.FASES] +
                  [("Total", lambda ref: f'=COUNTIF({rc},{ref})', "0")], 34)
    configurar_impresion(rs, "Causas de no cumplimiento", "portrait", area=f"A1:E{6 + len(causas)}")

    hoja_instrucciones(
        lib,
        "Planificar las próximas 3 semanas de trabajo con IWP tomados del backlog, confirmar recursos críticos y medir el cumplimiento del plan semanal (PPC) con sus causas de no cumplimiento.",
        "Cada semana, en la reunión de lookahead (jueves). Durante la superposición de fases, se usa un solo lookahead conjunto.",
        "Los superintendentes seleccionan los IWP; el Líder de WFP y los planificadores preparan la hoja; el Gerente de Construcción la aprueba.",
        [
            "Escriba el lunes de la semana 1 (B3); las fechas de los días se actualizan solas.",
            "Borre las filas amarillas de ejemplo.",
            "Registre los IWP seleccionados del backlog con sus restricciones abiertas y recursos críticos.",
            "Marque con «X» los días de ejecución de cada IWP en las 3 semanas.",
            "Un día marcado en rojo indica un IWP programado con restricciones abiertas: resuélvalo o sáquelo del plan.",
            "Al cerrar la semana 1, indique si cada IWP planificado se completó y, si no, la causa. El PPC se calcula en E3.",
            "Revise la hoja «Causas de no cumplimiento» y lleve las causas repetidas al registro de lecciones aprendidas.",
        ],
        [("Lookahead", [col for col in cols if not col.clave.startswith("d")] +
          [("S1 L … S3 S", "Días de las 3 semanas (lunes a sábado). Marque «X» los días de ejecución del IWP.", "Lista")])],
        justificacion="El procedimiento 3.0 y la comparación AWP–Lean Construction establecen el three-week look-ahead y el PPC como el nexo entre el backlog de IWP y la ejecución semanal.",
    )
    lib.crear_listas(extra={"Causas": causas})
    lib.guardar()


# ==========================================================================
# 8. Checklist de liberación de IWP
# ==========================================================================

CRITERIOS_CHECKLIST = [
    ("Documentos", "Planos y documentos IFC vigentes incluidos en el IWP", "Líder de Ingeniería"),
    ("Documentos", "Sin RFI abiertas que afecten al alcance del IWP", "Líder de Ingeniería"),
    ("Documentos", "Vistas 3D, isométricos o croquis del alcance incluidos", "Planificador de frente de trabajo"),
    ("Documentos", "Cantidades y HH estimadas verificadas (300 a 600 HH)", "Planificador de frente de trabajo"),
    ("Materiales", "100 % de materiales recibidos en almacén", "Gestor de materiales"),
    ("Materiales", "Materiales inspeccionados y conformes", "Líder de calidad (QA/QC)"),
    ("Materiales", "Materiales reservados para el IWP (y preparados por paquete si aplica)", "Gestor de materiales"),
    ("Materiales", "Documentación del proveedor disponible (planos certificados, manuales)", "Líder de Procura"),
    ("Recursos", "Cuadrilla y capataz asignados y acreditados", "Superintendente / capataz general"),
    ("Recursos", "Equipos y grúas confirmados para las fechas", "Gerente de Construcción"),
    ("Recursos", "Andamios montados e inspeccionados (o solicitados con fecha confirmada)", "Superintendente / capataz general"),
    ("Recursos", "Herramientas y consumibles disponibles", "Superintendente / capataz general"),
    ("Frente de trabajo", "Trabajos predecesores terminados y aceptados por calidad", "Superintendente / capataz general"),
    ("Frente de trabajo", "Frente de trabajo libre de interferencias con otras cuadrillas o fases", "Gerente de Construcción"),
    ("HSE", "JHA elaborado y aprobado", "Líder de HSE"),
    ("HSE", "Permisos de trabajo identificados (excavación, altura, caliente, izaje)", "Líder de HSE"),
    ("HSE", "Plan de izaje aprobado (si aplica)", "Líder de HSE"),
    ("Calidad", "ITP y formatos de inspección incluidos", "Líder de calidad (QA/QC)"),
    ("Calidad", "Procedimientos de trabajo aprobados (soldadura, concreto, pruebas)", "Líder de calidad (QA/QC)"),
    ("Registro", "Todas las restricciones del IWP liberadas o canceladas en el registro", "Líder de WFP"),
    ("Registro", "IWP incluido en el cronograma de nivel 5 y con sistema asignado", "Controles del proyecto"),
    ("Aprobación", "Revisado por el planificador de frente de trabajo", "Planificador de frente de trabajo"),
    ("Aprobación", "Revisado por el superintendente", "Superintendente / capataz general"),
    ("Aprobación", "Aprobado para liberación por el Líder de WFP", "Líder de WFP"),
]


def checklist_liberacion():
    lib = Libro(SALIDA / "Checklist_Liberacion_IWP.xlsx", "Checklist de liberación de IWP", ["Fases", "SiNo", "Roles"])
    cumple_ej = ["Sí"] * len(CRITERIOS_CHECKLIST)
    aplica_ej = ["Sí"] * len(CRITERIOS_CHECKLIST)
    aplica_ej[16] = "No"  # plan de izaje no aplica a zapatas
    cumple_ej[16] = "No aplica"
    filas = []
    for n, (cat, crit, rol) in enumerate(CRITERIOS_CHECKLIST):
        filas.append(dict(n=n + 1, cat=cat, crit=crit, aplica=aplica_ej[n], cumple=cumple_ej[n], rol=rol,
                          evid="Ver registro" if cat == "Registro" else "", obs=""))
    cols = [
        Col("n", "N.°", 5, "num", desc="Número del criterio."),
        Col("cat", "Categoría", 15, desc="Grupo del criterio."),
        Col("crit", "Criterio de liberación", 56, ajustar=True, desc="Condición que debe cumplirse antes de liberar el IWP."),
        Col("aplica", "Aplica", 8, lista="SiNo", desc="«No» si el criterio no aplica a este IWP."),
        Col("cumple", "Cumple", 11, lista_valores=["Sí", "No", "Pendiente", "No aplica"], desc="Sí, No, Pendiente o No aplica."),
        Col("evid", "Evidencia / referencia", 22, ajustar=True, desc="Documento, transmittal, vale de almacén, permiso, etc."),
        Col("rol", "Responsable (rol)", 24, lista="Roles", desc="Rol que confirma el criterio."),
        Col("obs", "Observaciones", 26, ajustar=True, desc="Comentarios."),
    ]
    ws, c, ult = hoja_datos(
        lib, "Checklist", "Checklist de liberación de IWP", cols, [], n_filas=len(CRITERIOS_CHECKLIST) + 6,
        params=[("iwp", "IWP:", "IWP-2.01-CIV-01-001", None), ("fase", "Fase:", "Fase 2", None), ("fecha", "Fecha:", D(2028, 1, 26), "fecha")],
        subtitulo=f"{dp.PROYECTO}. Ejemplo lleno para el IWP-2.01-CIV-01-001 (zapatas ejes A1–A4). Use una copia de esta hoja por IWP.",
        congelar_col=3, orientacion="portrait",
    )
    for n, datos in enumerate(filas):
        for col in cols:
            ws[f"{c(col.clave)}{FILA_DATOS + n}"] = datos.get(col.clave)
            ws[f"{c(col.clave)}{FILA_DATOS + n}"].fill = relleno("FFF2CC")
    p = ws._kit_celdas_param
    dvf = DataValidation(type="list", formula1="=Fases", allow_blank=True)
    ws.add_data_validation(dvf)
    dvf.add(p["fase"].replace("$", ""))
    ra, rc = rango(c, "aplica", ult), rango(c, "cumple", ult)
    ws["A4"] = "Planificador:"
    ws["B4"] = "(nombre)"
    ws["D4"] = "Cumplimiento:"
    ws["E4"] = f'=IFERROR(COUNTIFS({ra},"Sí",{rc},"Sí")/COUNTIF({ra},"Sí"),"")'
    ws["E4"].number_format = "0%"
    ws["G4"] = "Resultado:"
    ws["H4"] = f'=IF(COUNTIF({ra},"Sí")=0,"",IF(COUNTIFS({ra},"Sí",{rc},"<>Sí")=0,"LIBERAR","NO LIBERAR"))'
    for ref in ("A4", "D4", "G4"):
        ws[ref].font = Font(name=FUENTE, bold=True)
        ws[ref].alignment = Alignment(horizontal="right")
    for ref in ("B4", "E4", "H4"):
        ws[ref].font = Font(name=FUENTE, bold=True, size=11)
        ws[ref].border = BORDE
        ws[ref].fill = relleno(AZUL_CLARO)
    resaltar_fila(ws, "H4", 'H4="LIBERAR"', VERDE, VERDE_TEXTO, True)
    resaltar_fila(ws, "H4", 'H4="NO LIBERAR"', ROJO, ROJO_TEXTO, True)
    cu = f"${c('cumple')}{FILA_DATOS}"
    ap = f"${c('aplica')}{FILA_DATOS}"
    resaltar_fila(ws, f"{c('cumple')}{FILA_DATOS}:{c('cumple')}{ult}", f'AND({ap}="Sí",OR({cu}="No",{cu}="Pendiente",{cu}=""))', ROJO, ROJO_TEXTO)
    resaltar_fila(ws, f"{c('cumple')}{FILA_DATOS}:{c('cumple')}{ult}", f'{cu}="Sí"', VERDE, VERDE_TEXTO)

    # Registro de liberaciones
    regs = [
        ("IWP-1.01-CIV-01-001", "Fase 1", D(2027, 6, 11), 22, 22, "Líder de WFP", D(2027, 6, 11), "Liberado tras permiso de excavación (R-0001)."),
        ("IWP-1.02-TUB-01-001", "Fase 1", D(2027, 7, 19), 21, 21, "Líder de WFP", D(2027, 7, 19), ""),
        ("IWP-2.01-CIV-01-001", "Fase 2", D(2028, 1, 26), 23, 23, "Líder de WFP", D(2028, 1, 26), "Ver hoja «Checklist»."),
        ("IWP-2.01-EST-01-001", "Fase 2", D(2028, 4, 14), 24, 20, "", None, "No liberado: grúa y plan de izaje (R-0005, R-0008)."),
        ("IWP-2.03-ELE-01-001", "Fase 2", D(2028, 6, 14), 23, 21, "", None, "Pendiente andamio (R-0007) y documentación."),
        ("IWP-3.02-CIV-01-001", "Fase 3", D(2028, 6, 14), 22, 6, "", None, "Revisión preliminar; IWP en la ventana de la Fase 3 más adelante."),
    ]
    rcols = [
        COL_EJEMPLO,
        Col("iwp", "IWP", 21, desc="Código del IWP revisado."),
        Col("fase", "Fase del proyecto", 11, lista="Fases", desc="Fase del proyecto."),
        Col("cwp", "CWP", 17, formula=lambda f, c, p: f'=IF({c("iwp")}{f}="","","CWP"&MID({c("iwp")}{f},4,12))', desc="Calculado a partir del código del IWP."),
        Col("fecha", "Fecha de revisión", 11, "fecha", desc="Fecha en que se aplicó el checklist."),
        Col("aplic", "Criterios aplicables", 10, "num", desc="Cantidad de criterios con «Aplica = Sí»."),
        Col("cumpl", "Criterios cumplidos", 10, "num", desc="Cantidad de criterios aplicables con «Cumple = Sí»."),
        Col("pct", "% cumplimiento", 10, "pct", formula=lambda f, c, p: f'=IF(OR({c("aplic")}{f}="",{c("aplic")}{f}=0),"",{c("cumpl")}{f}/{c("aplic")}{f})', desc="Calculado."),
        Col("res", "Resultado", 12, formula=lambda f, c, p: f'=IF({c("pct")}{f}="","",IF({c("pct")}{f}=1,"LIBERAR","NO LIBERAR"))', desc="Calculado: LIBERAR solo con 100 % de cumplimiento."),
        Col("por", "Liberado por (rol)", 18, lista="Roles", desc="Rol que firmó la liberación."),
        Col("f_lib", "Fecha de liberación", 11, "fecha", desc="Fecha de liberación (vacía si no se liberó)."),
        Col("obs", "Observaciones", 36, ajustar=True, desc="Comentarios; restricciones pendientes."),
    ]
    rws, rc_, rult = hoja_datos(lib, "Registro de liberaciones", "Registro de liberaciones de IWP", rcols,
                                ejemplos(regs, ["iwp", "fase", "fecha", "aplic", "cumpl", "por", "f_lib", "obs"]), n_filas=300)
    r = f"${rc_('res')}{FILA_DATOS}"
    resaltar_fila(rws, f"{rc_('res')}{FILA_DATOS}:{rc_('res')}{rult}", f'{r}="LIBERAR"', VERDE, VERDE_TEXTO, True)
    resaltar_fila(rws, f"{rc_('res')}{FILA_DATOS}:{rc_('res')}{rult}", f'{r}="NO LIBERAR"', ROJO, ROJO_TEXTO, True)

    hoja_instrucciones(
        lib,
        "Verificar, criterio por criterio, que un IWP está libre de restricciones antes de entregarlo a campo, y dejar registro de cada liberación.",
        "Antes de pasar cada IWP al backlog (liberación). Se repite si el IWP vuelve de campo por una restricción no detectada.",
        "El planificador de frente de trabajo completa el checklist con los responsables de cada criterio; el superintendente revisa y el Líder de WFP aprueba.",
        [
            "Haga una copia de la hoja «Checklist» por cada IWP (clic derecho en la pestaña > Mover o copiar > Crear una copia).",
            "Escriba el código del IWP, la fase y la fecha en la fila 3, y el planificador en la fila 4.",
            "Para cada criterio, indique si aplica y si cumple, con la evidencia.",
            "El resultado (H4) es «LIBERAR» solo si todos los criterios aplicables cumplen.",
            "Registre el resultado en la hoja «Registro de liberaciones».",
            "Las filas amarillas son de ejemplo: bórrelas o reemplácelas.",
        ],
        [("Checklist", cols), ("Registro de liberaciones", rcols)],
        justificacion="Las guías de inicio rápido y el procedimiento 3.0 exigen que el IWP llegue a campo libre de restricciones; el checklist hace verificable esa regla y alimenta el KPI de IWP liberados sin restricciones.",
    )
    lib.crear_listas()
    lib.guardar()


# ==========================================================================
# 9. Tablero de KPI
# ==========================================================================

KPIS = [
    # código, nombre, fórmula, unidad, sentido, metas F1-F3, tipo
    ("K01", "EWP emitidos a tiempo", "EWP emitidos IFC a tiempo ÷ EWP con fecha requerida vencida", "%", "≥", (0.80, 0.90, 0.95), "Proceso"),
    ("K02", "Materiales disponibles al liberar", "IWP liberados con 100 % de materiales ÷ IWP liberados", "%", "≥", (0.90, 0.95, 0.98), "Proceso"),
    ("K03", "IWP liberados sin restricciones", "IWP liberados sin restricciones abiertas ÷ IWP liberados", "%", "≥", (0.85, 0.95, 0.98), "Proceso"),
    ("K04", "Backlog de IWP liberados", "HH de IWP liberados no iniciados ÷ HH ejecutadas por semana", "semanas", "≥", (2, 2, 3), "Proceso"),
    ("K05", "Restricciones liberadas a tiempo", "Restricciones liberadas a tiempo ÷ restricciones liberadas", "%", "≥", (0.80, 0.90, 0.95), "Proceso"),
    ("K06", "Atraso promedio de restricciones", "Suma de días de atraso ÷ N.° de restricciones vencidas", "días", "≤", (7, 5, 3), "Proceso"),
    ("K07", "Cobertura de planificadores", "Trabajadores directos ÷ planificadores", "trab./planif.", "≤", (55, 50, 50), "Proceso"),
    ("K08", "Porcentaje de plan cumplido (PPC)", "IWP completados según plan ÷ IWP planificados", "%", "≥", (0.75, 0.80, 0.85), "Resultado"),
    ("K09", "IWP devueltos", "IWP devueltos de campo ÷ IWP entregados a campo", "%", "≤", (0.08, 0.05, 0.03), "Resultado"),
    ("K10", "Factor de productividad", "HH ganadas ÷ HH gastadas", "índice", "≥", (0.95, 1.00, 1.05), "Resultado"),
    ("K11", "Tiempo productivo (tool time)", "Tiempo en trabajo directo ÷ tiempo observado", "%", "≥", (0.40, 0.44, 0.46), "Resultado"),
    ("K12", "Retrabajo", "HH de retrabajo ÷ HH gastadas", "%", "≤", (0.04, 0.03, 0.02), "Resultado"),
    ("K13", "SPI de construcción", "Valor ganado ÷ valor planificado", "índice", "≥", (0.95, 0.98, 1.00), "Resultado"),
]
FORMATO_UNIDAD = {"%": "0.0%", "semanas": "0.0", "días": "0.0", "trab./planif.": "0", "índice": "0.00"}


def tablero_kpi():
    lib = Libro(SALIDA / "Tablero_KPI.xlsx", "Tablero de KPI por fase", ["Fases"])
    # --- Datos semanales
    campos = [
        ("iwp_plan", "IWP planificados en la semana", "num"), ("iwp_comp", "IWP completados según plan", "num"),
        ("iwp_lib", "IWP liberados en la semana", "num"), ("iwp_sin", "IWP liberados sin restricciones", "num"),
        ("iwp_mat", "IWP liberados con 100 % de materiales", "num"), ("iwp_ent", "IWP entregados a campo", "num"),
        ("iwp_dev", "IWP devueltos de campo", "num"), ("hh_back", "HH de IWP en backlog (fin de semana)", "num"),
        ("hh_ejec", "HH ejecutadas en la semana", "num"), ("r_lib", "Restricciones liberadas", "num"),
        ("r_tiempo", "Restricciones liberadas a tiempo", "num"), ("r_venc", "Restricciones vencidas (al corte)", "num"),
        ("r_dias", "Suma de días de atraso de vencidas", "num"), ("ewp_venc", "EWP con fecha requerida vencida", "num"),
        ("ewp_ok", "EWP emitidos a tiempo", "num"), ("hh_gan", "HH ganadas", "num"), ("hh_gas", "HH gastadas", "num"),
        ("hh_ret", "HH de retrabajo", "num"), ("trab", "Trabajadores directos", "num"), ("planif", "Planificadores de frente de trabajo", "num"),
        ("tool", "Tool time medido (si hubo estudio)", "pct"), ("ve", "Valor ganado", "num"), ("vp", "Valor planificado", "num"),
    ]
    semanas = [
        ("Fase 1", D(2027, 9, 6), (40, 29, 38, 31, 34, 39, 4, 9500, 5200, 22, 16, 6, 54, 5, 4, 5000, 5400, 240, 205, 4, 0.39, 5100, 5400)),
        ("Fase 1", D(2027, 9, 13), (42, 32, 40, 34, 37, 41, 3, 10400, 5300, 25, 20, 5, 38, 4, 3, 5150, 5350, 210, 210, 4, None, 5250, 5450)),
        ("Fase 1", D(2027, 9, 20), (41, 33, 42, 37, 39, 40, 3, 11200, 5400, 24, 20, 4, 26, 3, 3, 5300, 5380, 180, 215, 4, None, 5380, 5500)),
        ("Fase 2", D(2028, 5, 29), (150, 118, 152, 142, 145, 148, 8, 52000, 19500, 80, 70, 9, 52, 12, 11, 19300, 19600, 610, 790, 16, 0.43, 19300, 19900)),
        ("Fase 2", D(2028, 6, 5), (155, 125, 150, 143, 146, 152, 7, 55000, 20100, 84, 75, 8, 41, 10, 9, 20200, 20100, 580, 810, 16, None, 20150, 20500)),
        ("Fase 2", D(2028, 6, 12), (158, 130, 160, 153, 156, 156, 6, 58000, 20800, 86, 78, 8, 36, 9, 9, 21000, 20600, 560, 830, 17, None, 20950, 21200)),
        ("Fase 3", D(2029, 3, 5), (70, 60, 72, 71, 71, 70, 2, 30000, 8900, 38, 37, 2, 5, 6, 6, 9300, 8900, 170, 360, 8, 0.46, 9150, 9100)),
        ("Fase 3", D(2029, 3, 12), (72, 62, 74, 73, 73, 72, 2, 31000, 9100, 40, 38, 2, 6, 5, 5, 9550, 9100, 160, 370, 8, None, 9400, 9350)),
        ("Fase 3", D(2029, 3, 19), (73, 63, 75, 74, 74, 73, 1, 32500, 9300, 41, 40, 1, 2, 5, 5, 9800, 9250, 150, 375, 8, None, 9650, 9600)),
    ]
    filas = []
    for fase, sem, valores in semanas:
        f = {"ejemplo": EJ, "semana": sem, "fase": fase}
        f.update({k: v for (k, _, _), v in zip(campos, valores) if v is not None})
        filas.append(f)

    def div(a, b, fmt_pct=True):
        return lambda f, c, p: f'=IF(OR({c(b)}{f}="",{c(b)}{f}=0),"",{c(a)}{f}/{c(b)}{f})'

    cols = [COL_EJEMPLO, Col("semana", "Semana (lunes)", 11, "fecha", desc="Lunes de la semana reportada."),
            Col("fase", "Fase del proyecto", 10, lista="Fases", desc="Fase del proyecto.")]
    for k, t, tipo in campos:
        cols.append(Col(k, t, 11, tipo, desc=f"Dato semanal: {t.lower()}."))
    cols += [
        Col("ppc", "PPC semanal", 9, "pct", formula=div("iwp_comp", "iwp_plan"), desc="Calculado: IWP completados ÷ IWP planificados."),
        Col("sinr", "% IWP sin restricciones", 10, "pct", formula=div("iwp_sin", "iwp_lib"), desc="Calculado: IWP liberados sin restricciones ÷ IWP liberados."),
        Col("back", "Backlog (semanas)", 9, "dec", formula=div("hh_back", "hh_ejec"), desc="Calculado: HH en backlog ÷ HH ejecutadas en la semana."),
        Col("fp", "Factor de productividad", 10, "dec", formula=div("hh_gan", "hh_gas"), desc="Calculado: HH ganadas ÷ HH gastadas."),
    ]
    ws, c, ult = hoja_datos(lib, "Datos semanales", "Datos semanales para el cálculo de KPI", cols, filas, n_filas=300, congelar_col=3, alto_enc=58)
    for k in ("ppc", "sinr"):
        pass

    # --- Metas
    ms = lib.wb.create_sheet("Metas")
    estilo_titulo(ms, "Metas de KPI por fase del proyecto", "Puede ajustar las metas; el tablero se recalcula. Sentido «≥»: mayor es mejor; «≤»: menor es mejor.", 9)
    enc = ["Código", "KPI", "Fórmula", "Unidad", "Sentido", "Meta Fase 1", "Meta Fase 2", "Meta Fase 3", "Tipo"]
    for i, t in enumerate(enc, 1):
        celda = ms.cell(row=5, column=i, value=t)
        celda.font = Font(name=FUENTE, bold=True, color=BLANCO)
        celda.fill = relleno(AZUL)
        celda.border = BORDE
        celda.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    for n, (cod, nom, form, uni, sent, metas, tipo) in enumerate(KPIS, 6):
        vals = [cod, nom, form, uni, sent, *metas, tipo]
        for i, v in enumerate(vals, 1):
            celda = ms.cell(row=n, column=i, value=v)
            celda.border = BORDE
            celda.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if i in (1, 4, 5, 9) else None)
            if 6 <= i <= 8:
                celda.number_format = FORMATO_UNIDAD[uni]
                celda.fill = relleno(AZUL_CLARO)
    for col, w in zip("ABCDEFGHI", (8, 30, 50, 12, 8, 11, 11, 11, 11)):
        ms.column_dimensions[col].width = w
    dv = DataValidation(type="list", formula1='"≥,≤"', allow_blank=False)
    ms.add_data_validation(dv)
    dv.add(f"E6:E{5 + len(KPIS)}")
    ms.freeze_panes = "C6"
    configurar_impresion(ms, "Metas de KPI", "landscape", "5:5", f"A1:I{5 + len(KPIS)}")

    # --- Tablero
    ts = lib.wb.create_sheet("Tablero")
    estilo_titulo(ts, "Tablero de KPI por fase del proyecto", f"{dp.PROYECTO}. Valores acumulados de la hoja «Datos semanales». Verde: cumple la meta; ámbar: a menos de 10 % de la meta; rojo: más de 10 % fuera de la meta.", 16)
    fila_enc = 5
    enc1 = ["Código", "KPI", "Tipo", "Unidad", "Sentido"]
    for i, t in enumerate(enc1, 1):
        ts.cell(row=fila_enc, column=i, value=t)
        ts.merge_cells(start_row=fila_enc, start_column=i, end_row=fila_enc + 1, end_column=i)
    col0 = len(enc1) + 1
    for n, fase in enumerate(dp.FASES):
        ci = col0 + n * 3
        ts.cell(row=fila_enc, column=ci, value=f"{fase} – {dp.NOMBRE_FASE[fase]}")
        ts.merge_cells(start_row=fila_enc, start_column=ci, end_row=fila_enc, end_column=ci + 2)
        for k, t in enumerate(("Valor", "Meta", "Semáforo")):
            ts.cell(row=fila_enc + 1, column=ci + k, value=t)
    ts.cell(row=fila_enc, column=col0 + 9, value="Tendencia F1 → F3")
    ts.merge_cells(start_row=fila_enc, start_column=col0 + 9, end_row=fila_enc + 1, end_column=col0 + 9)
    for fila in (fila_enc, fila_enc + 1):
        for i in range(1, col0 + 10):
            celda = ts.cell(row=fila, column=i)
            celda.font = Font(name=FUENTE, bold=True, color=BLANCO, size=10)
            celda.fill = relleno(AZUL)
            celda.border = BORDE
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    D_ = {k: f"'Datos semanales'!{rango(c, k, ult)}" for k in [x for x, _, _ in campos] + ["fase", "back"]}

    def suma(k, fref):
        return f'SUMIFS({D_[k]},{D_["fase"]},{fref})'

    valor = {
        "K01": lambda f: f"{suma('ewp_ok', f)}/{suma('ewp_venc', f)}",
        "K02": lambda f: f"{suma('iwp_mat', f)}/{suma('iwp_lib', f)}",
        "K03": lambda f: f"{suma('iwp_sin', f)}/{suma('iwp_lib', f)}",
        "K04": lambda f: f'AVERAGEIFS({D_["back"]},{D_["fase"]},{f})',
        "K05": lambda f: f"{suma('r_tiempo', f)}/{suma('r_lib', f)}",
        "K06": lambda f: f"{suma('r_dias', f)}/{suma('r_venc', f)}",
        "K07": lambda f: f"{suma('trab', f)}/{suma('planif', f)}",
        "K08": lambda f: f"{suma('iwp_comp', f)}/{suma('iwp_plan', f)}",
        "K09": lambda f: f"{suma('iwp_dev', f)}/{suma('iwp_ent', f)}",
        "K10": lambda f: f"{suma('hh_gan', f)}/{suma('hh_gas', f)}",
        "K11": lambda f: f'AVERAGEIFS({D_["tool"]},{D_["fase"]},{f},{D_["tool"]},"<>")',
        "K12": lambda f: f"{suma('hh_ret', f)}/{suma('hh_gas', f)}",
        "K13": lambda f: f"{suma('ve', f)}/{suma('vp', f)}",
    }
    primera = fila_enc + 2
    for n, (cod, nom, form, uni, sent, metas, tipo) in enumerate(KPIS):
        fila = primera + n
        mfila = 6 + n
        for i, v in enumerate((cod, nom, tipo, uni, f"=Metas!E{mfila}"), 1):
            ts.cell(row=fila, column=i, value=v)
        for k, fase in enumerate(dp.FASES):
            ci = col0 + k * 3
            fref = f'"{fase}"'
            v = ts.cell(row=fila, column=ci, value=f'=IFERROR({valor[cod](fref)},"")')
            m = ts.cell(row=fila, column=ci + 1, value=f"=Metas!{get_column_letter(6 + k)}{mfila}")
            vl, ml, sl = get_column_letter(ci), get_column_letter(ci + 1), f"$E{fila}"
            ts.cell(row=fila, column=ci + 2, value=(
                f'=IF({vl}{fila}="","",IF({sl}="≥",IF({vl}{fila}>={ml}{fila},"Verde",IF({vl}{fila}>={ml}{fila}*0.9,"Ámbar","Rojo")),'
                f'IF({vl}{fila}<={ml}{fila},"Verde",IF({vl}{fila}<={ml}{fila}*1.1,"Ámbar","Rojo"))))'))
            v.number_format = m.number_format = FORMATO_UNIDAD[uni]
        v1, v3 = get_column_letter(col0), get_column_letter(col0 + 6)
        ts.cell(row=fila, column=col0 + 9, value=(
            f'=IF(OR({v1}{fila}="",{v3}{fila}=""),"",IF({v3}{fila}={v1}{fila},"Sin cambio",'
            f'IF(({v3}{fila}>{v1}{fila})=($E{fila}="≥"),"Mejora","Empeora")))'))
        for i in range(1, col0 + 10):
            celda = ts.cell(row=fila, column=i)
            celda.border = BORDE
            celda.font = Font(name=FUENTE, size=10, bold=(i == 1))
            celda.alignment = Alignment(horizontal="left" if i == 2 else "center", vertical="center", wrap_text=(i == 2))
        ts.row_dimensions[fila].height = 22
    ultima_t = primera + len(KPIS) - 1
    for k in range(3):
        col = get_column_letter(col0 + k * 3 + 2)
        semaforo_texto(ts, f"{col}{primera}:{col}{ultima_t}", f"{col}{primera}")
    tcol = get_column_letter(col0 + 9)
    resaltar_fila(ts, f"{tcol}{primera}:{tcol}{ultima_t}", f'{tcol}{primera}="Mejora"', VERDE, VERDE_TEXTO)
    resaltar_fila(ts, f"{tcol}{primera}:{tcol}{ultima_t}", f'{tcol}{primera}="Empeora"', ROJO, ROJO_TEXTO)
    anchos = {1: 7, 2: 32, 3: 10, 4: 10, 5: 7}
    for i in range(1, col0 + 10):
        ts.column_dimensions[get_column_letter(i)].width = anchos.get(i, 10)
    ts.freeze_panes = f"C{primera}"

    # Gráfico: KPI porcentuales por fase
    ch = BarChart()
    ch.type = "col"
    ch.grouping = "clustered"
    ch.title = "KPI porcentuales por fase del proyecto"
    ch.y_axis.title = "Valor"
    ch.y_axis.numFmt = "0%"
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 1
    ch.height, ch.width = 8.5, 22
    filas_pct = [primera + i for i, k in enumerate(KPIS) if k[3] == "%" and k[4] == "≥"]
    # Tabla auxiliar para el gráfico (debajo del tablero)
    aux = ultima_t + 3
    ts.cell(row=aux, column=1, value="Datos del gráfico").font = Font(name=FUENTE, bold=True, color=AZUL)
    ts.cell(row=aux + 1, column=2, value="KPI")
    for k, fase in enumerate(dp.FASES):
        ts.cell(row=aux + 1, column=3 + k, value=fase)
    for n, fr in enumerate(filas_pct):
        ts.cell(row=aux + 2 + n, column=2, value=f"=A{fr}&\" \"&B{fr}")
        for k in range(3):
            celda = ts.cell(row=aux + 2 + n, column=3 + k, value=f"=IF({get_column_letter(col0 + k * 3)}{fr}=\"\",0,{get_column_letter(col0 + k * 3)}{fr})")
            celda.number_format = "0%"
    datos_ref = Reference(ts, min_col=3, max_col=5, min_row=aux + 1, max_row=aux + 1 + len(filas_pct))
    cats = Reference(ts, min_col=2, min_row=aux + 2, max_row=aux + 1 + len(filas_pct))
    ch.add_data(datos_ref, titles_from_data=True)
    ch.set_categories(cats)
    ts.add_chart(ch, f"B{aux + 3 + len(filas_pct)}")
    configurar_impresion(ts, "Tablero de KPI", "landscape", f"{fila_enc}:{fila_enc + 1}", f"A1:{tcol}{aux + 22 + len(filas_pct)}")
    ts.page_setup.fitToHeight = 1
    lib.wb.move_sheet("Tablero", offset=-2)
    lib.wb.move_sheet("Metas", offset=-1)

    hoja_instrucciones(
        lib,
        "Medir con las mismas fórmulas, en las tres fases del proyecto, si AWP se está implementando (KPI de proceso) y si mejora la productividad (KPI de resultado), y compararlas contra metas progresivas por fase.",
        "Los datos se cargan cada semana; el tablero se revisa en el reporte semanal y en el Comité AWP mensual. Al cierre de cada fase se usa para el informe de cierre.",
        "Controles del proyecto carga los datos con la información del Líder de WFP, ingeniería y procura; el AWP Champion analiza y presenta el tablero.",
        [
            "Borre las filas amarillas de ejemplo en «Datos semanales».",
            "Cargue una fila por semana y por fase del proyecto (si dos fases están en construcción, dos filas por semana).",
            "Ingrese solo datos base; los KPI se calculan en el «Tablero».",
            "El tool time se registra solo en las semanas en que se hizo un estudio de muestreo.",
            "Revise o ajuste las metas en la hoja «Metas».",
            "Use la columna «Tendencia F1 → F3» para mostrar el efecto de las lecciones aprendidas entre fases.",
        ],
        [("Datos semanales", cols),
         ("Tablero", [("Valor", "KPI acumulado de la fase (fórmula de la hoja «Metas»).", "Calculado"),
                      ("Meta", "Meta de la fase (de la hoja «Metas»).", "Calculado"),
                      ("Semáforo", "Verde / Ámbar / Rojo según la meta y el sentido.", "Calculado"),
                      ("Tendencia F1 → F3", "Mejora o empeora el valor de la Fase 3 respecto de la Fase 1.", "Calculado")])],
        justificacion="El curso 2023, Omega 365 y el marco educativo del CII piden medir AWP con KPI, pero sus listas se perdieron en la conversión. Este tablero reúne los KPI del plan con fórmulas únicas, requisito para comparar fases.",
    )
    lib.crear_listas()
    lib.guardar()


# ==========================================================================
# 10. Registro de riesgos
# ==========================================================================

RIESGOS = [
    ("R01", "Todas las fases", "Planificación temprana (FEL)", "Organización y liderazgo", "Falta de patrocinio del cliente y de la gerencia; AWP percibido como «papeleo adicional».", "Beneficios de AWP no comunicados; presión por avance inmediato.", "Abandono parcial del proceso; KPI sin seguimiento.", 3, 5, "Plan firmado por el cliente; Comité AWP mensual con KPI; comunicar resultados tempranos de la Fase 1.", "Gerente del proyecto del cliente", D(2028, 7, 1), "En tratamiento", 2, 4),
    ("R02", "Fase 2", "Ingeniería", "Procesos", "Ingeniería no entrega EWP en la secuencia del PoC.", "Planificación de ingeniería por disciplina y no por CWP.", "CWP sin planos IFC; backlog insuficiente.", 4, 4, "Fechas de EWP derivadas del PoC en el contrato; KPI K01; revisión quincenal del plan de liberación.", "Líder de Ingeniería", D(2028, 7, 1), "En tratamiento", 3, 3),
    ("R03", "Fase 2", "Construcción", "Personas y competencias", "Falta de planificadores de frente de trabajo con experiencia (y rotación de roles clave).", "Mercado laboral competido; perfil poco conocido.", "IWP de baja calidad; restricciones no detectadas.", 4, 4, "Formar planificadores en la Fase 1 y transferirlos; reclutamiento anticipado; formación de 40 h; suplentes para Champion y Líder de WFP.", "AWP Champion", D(2028, 7, 1), "En tratamiento", 2, 3),
    ("R04", "Fases 1 y 2", "Procura", "Herramientas e información", "Procura no sigue los materiales por CWP/IWP.", "Órdenes de compra sin código de CWP.", "No se sabe qué IWP tiene material completo.", 3, 4, "PWP por CWP; código de CWP en órdenes de compra y almacén; reporte semanal de disponibilidad.", "Líder de Procura", D(2028, 7, 1), "En tratamiento", 2, 3),
    ("R05", "Fases 2 y 3", "Planificación temprana (FEL)", "Contratos", "Subcontratistas sin conocimiento ni obligación contractual de AWP.", "Requisitos AWP no incluidos en los subcontratos.", "Frentes fuera del proceso; datos incompletos.", 3, 4, "Cláusulas AWP en subcontratos; formación de inicio; planificadores del subcontratista integrados.", "Gerente de Proyecto", D(2028, 7, 1), "Abierto", 2, 3),
    ("R06", "Todas las fases", "Construcción", "Procesos", "Se liberan IWP con restricciones abiertas por presión de avance.", "Backlog corto y metas de avance semanales.", "IWP devueltos; tiempo improductivo.", 3, 4, "Checklist de liberación obligatorio; KPI K03; auditorías.", "Líder de WFP", D(2028, 7, 1), "En tratamiento", 2, 3),
    ("R07", "Fase 2", "Planificación temprana (FEL)", "Herramientas e información", "Herramientas y datos desintegrados (modelo 3D, cronograma y materiales).", "Sistemas de distintos proveedores sin codificación común.", "Retrabajo manual; errores en IWP.", 3, 3, "Coordinador de IM desde FEL; codificación única; integración probada antes de la construcción.", "Coordinador de gestión de información", D(2028, 7, 1), "En tratamiento", 2, 2),
    ("R08", "Fase 2", "Construcción", "Procesos", "Backlog insuficiente por avance de ingeniería o procura por detrás del PoC.", "Atrasos de EWP y PWP.", "Cuadrillas sin trabajo liberado; baja productividad.", 3, 4, "Seguimiento del backlog proyectado a 8 semanas; alertas tempranas; trabajo de «plan B».", "Gerente de Construcción", D(2028, 7, 1), "Materializado", 3, 3),
    ("R09", "Fases 2 y 3", "Construcción", "Interfaces entre fases", "Conflictos de recursos e interfaces entre fases superpuestas.", "Grúas, accesos y almacenes compartidos.", "Interferencias y demoras en ambas fases.", 4, 3, "Registro único de restricciones; lookahead conjunto; tablero de recursos compartidos.", "Gerente de Construcción", D(2028, 7, 1), "Abierto", 2, 3),
    ("R10", "Fases 1 y 2", "Comisionamiento", "Medición", "Lecciones de la Fase 1 no se aplican en la Fase 2.", "Falta de tiempo y de responsable de cada lección.", "Se repiten errores; KPI no mejoran entre fases.", 3, 3, "Puerta de control con revisión de lecciones; cada lección con responsable y fecha; KPI con fórmulas únicas.", "AWP Champion", D(2028, 7, 1), "Cerrado", 1, 3),
]
CLAVES_RIESGO = ["id", "fase", "etapa", "cat", "riesgo", "causa", "consec", "p", "i", "mitig", "rol", "f_rev", "estado", "pr", "ir"]


def registro_riesgos():
    lib = Libro(SALIDA / "Registro_Riesgos.xlsx", "Registro de riesgos de implementación AWP", ["FasesExt", "Etapas", "Roles"])

    def nivel(a, b):
        return lambda f, c, p: f'=IF(OR({c(a)}{f}="",{c(b)}{f}=""),"",{c(a)}{f}*{c(b)}{f})'

    def clasif(n):
        return lambda f, c, p: f'=IF({c(n)}{f}="","",IF({c(n)}{f}>=12,"Alto",IF({c(n)}{f}>=6,"Medio","Bajo")))'

    cols = [
        COL_EJEMPLO,
        Col("id", "ID", 6, desc="Identificador del riesgo (R01, R02…)."),
        Col("fase", "Fase del proyecto", 13, lista="FasesExt", desc="Fase o fases del proyecto afectadas."),
        Col("etapa", "Etapa del ciclo de vida", 16, lista="Etapas", desc="Etapa en que el riesgo se origina o se gestiona."),
        Col("cat", "Categoría", 16, lista="CategoriasRiesgo", desc="Categoría del riesgo de implementación."),
        Col("riesgo", "Descripción del riesgo", 36, ajustar=True, desc="Evento incierto que afectaría la implementación de AWP."),
        Col("causa", "Causa", 26, ajustar=True, desc="Por qué podría ocurrir."),
        Col("consec", "Consecuencia", 26, ajustar=True, desc="Qué pasaría si ocurre."),
        Col("p", "Probabilidad (1-5)", 9, "num", lista="Niveles", desc="1 muy baja (< 10 %) … 5 muy alta (> 70 %)."),
        Col("i", "Impacto (1-5)", 9, "num", lista="Niveles", desc="1 sin efecto visible en los KPI … 5 abandono de AWP o impacto en la ruta crítica."),
        Col("nivel", "Nivel (P×I)", 8, "num", formula=nivel("p", "i"), desc="Calculado: probabilidad × impacto."),
        Col("clas", "Clasificación", 10, formula=clasif("nivel"), desc="Calculado: Alto (12–25), Medio (6–11), Bajo (1–5)."),
        Col("mitig", "Mitigación", 40, ajustar=True, desc="Acciones para reducir probabilidad o impacto."),
        Col("rol", "Responsable (rol)", 22, lista="Roles", desc="Rol dueño del riesgo."),
        Col("f_rev", "Próxima revisión", 11, "fecha", desc="Fecha de la próxima revisión (Comité AWP)."),
        Col("estado", "Estado", 13, lista="EstadosRiesgo", desc="Abierto, En tratamiento, Cerrado o Materializado."),
        Col("pr", "P residual", 8, "num", lista="Niveles", desc="Probabilidad después de la mitigación."),
        Col("ir", "I residual", 8, "num", lista="Niveles", desc="Impacto después de la mitigación."),
        Col("nr", "Nivel residual", 8, "num", formula=nivel("pr", "ir"), desc="Calculado."),
        Col("cr", "Clasificación residual", 11, formula=clasif("nr"), desc="Calculado."),
        Col("obs", "Observaciones", 26, ajustar=True, desc="Comentarios."),
    ]
    ws, c, ult = hoja_datos(lib, "Riesgos", "Registro de riesgos de implementación AWP", cols, ejemplos(RIESGOS, CLAVES_RIESGO), n_filas=150, congelar_col=2)
    for k in ("clas", "cr"):
        rng = f"{c(k)}{FILA_DATOS}:{c(k)}{ult}"
        p = f"{c(k)}{FILA_DATOS}"
        resaltar_fila(ws, rng, f'{p}="Alto"', ROJO, ROJO_TEXTO, True)
        resaltar_fila(ws, rng, f'{p}="Medio"', AMBAR, "7F6000", True)
        resaltar_fila(ws, rng, f'{p}="Bajo"', VERDE, VERDE_TEXTO, True)
    e = f"${c('estado')}{FILA_DATOS}"
    resaltar_fila(ws, f"{c('estado')}{FILA_DATOS}:{c('estado')}{ult}", f'{e}="Materializado"', ROJO, ROJO_TEXTO, True)

    # Mapa de calor
    ms = lib.wb.create_sheet("Mapa de calor")
    estilo_titulo(ms, "Mapa de calor de riesgos (inherentes, no cerrados)", "Cantidad de riesgos por probabilidad e impacto. Calculado a partir de la hoja «Riesgos».", 7)
    rp, ri_, re_ = (f"Riesgos!{rango(c, k, ult)}" for k in ("p", "i", "estado"))
    ms.cell(row=4, column=1, value="Probabilidad \\ Impacto")
    for i in range(1, 6):
        ms.cell(row=4, column=1 + i, value=i)
    for n, prob in enumerate(range(5, 0, -1)):
        fila = 5 + n
        ms.cell(row=fila, column=1, value=prob)
        for imp in range(1, 6):
            celda = ms.cell(row=fila, column=1 + imp, value=f'=COUNTIFS({rp},{prob},{ri_},{imp},{re_},"<>Cerrado")')
            nivel_ = prob * imp
            celda.fill = relleno(ROJO if nivel_ >= 12 else AMBAR if nivel_ >= 6 else VERDE)
            celda.alignment = Alignment(horizontal="center", vertical="center")
            celda.font = Font(name=FUENTE, bold=True, size=14)
            celda.border = BORDE
        ms.row_dimensions[fila].height = 36
    for i in range(1, 7):
        ms.cell(row=4, column=i).font = Font(name=FUENTE, bold=True, color=BLANCO)
        ms.cell(row=4, column=i).fill = relleno(AZUL)
        ms.cell(row=4, column=i).alignment = Alignment(horizontal="center", wrap_text=True)
        ms.column_dimensions[get_column_letter(i)].width = 12
    for fila in range(5, 10):
        ms.cell(row=fila, column=1).font = Font(name=FUENTE, bold=True, color=BLANCO)
        ms.cell(row=fila, column=1).fill = relleno(AZUL)
        ms.cell(row=fila, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ms.column_dimensions["A"].width = 16
    ms["A11"] = "Verde: bajo (1–5) · Ámbar: medio (6–11) · Rojo: alto (12–25)"
    ms["A11"].font = Font(name=FUENTE, italic=True, size=9)
    tabla_resumen(ms, 13, 1, "Riesgos por fase y clasificación (no cerrados)", dp.FASES_EXT,
                  [(k, (lambda ref, k=k: f'=COUNTIFS(Riesgos!{rango(c, "fase", ult)},{ref},Riesgos!{rango(c, "clas", ult)},"{k}",{re_},"<>Cerrado")'), "0")
                   for k in ("Alto", "Medio", "Bajo")], 16)
    configurar_impresion(ms, "Mapa de calor", "portrait", area="A1:F22")

    hoja_instrucciones(
        lib,
        "Identificar, valorar y dar seguimiento a los riesgos de implementar AWP (no a los riesgos técnicos del proyecto), con su mitigación y responsable.",
        "Se elabora en la planificación temprana global y se revisa en el Comité AWP mensual y al inicio de cada fase del proyecto.",
        "El AWP Champion mantiene el registro; cada responsable actualiza el estado de sus riesgos.",
        [
            "Borre las filas amarillas de ejemplo o ajústelas a su proyecto.",
            "Describa el riesgo, su causa y su consecuencia.",
            "Valore probabilidad e impacto de 1 a 5 (ver escala en la sección «Riesgos de implementación» del plan).",
            "Defina la mitigación, el responsable y la fecha de revisión.",
            "Después de mitigar, valore el riesgo residual.",
            "Revise el «Mapa de calor» en el Comité AWP.",
        ],
        [("Riesgos", cols)],
        justificacion="El plan de implementación identifica riesgos propios de AWP (patrocinio, planificadores, secuencia de ingeniería, interfaces entre fases) que requieren seguimiento formal durante las tres fases.",
    )
    lib.crear_listas(extra={"CategoriasRiesgo": dp.CATEGORIAS_RIESGO, "EstadosRiesgo": dp.ESTADOS_RIESGO, "Niveles": dp.NIVELES_1_5})
    lib.guardar()


# ==========================================================================
# 11. Registro de lecciones aprendidas
# ==========================================================================

LECCIONES = [
    ("LA-01", "Fase 1", "Construcción", "Definición de paquetes", "Negativa", "Los IWP de movimiento de tierras definidos por volumen total duraban 3 semanas; el PPC no reflejaba los problemas reales.", "No se aplicó el criterio de una semana por IWP.", "Medio", "Limitar los IWP a 1 semana y 600 HH.", "Incluir la regla en la plantilla de IWP de la Fase 2.", "Plantilla_IWP.docx", "Fase 2", "Líder de WFP", D(2027, 12, 15), "Verificada", "PPC de Fase 2 con IWP de ≤ 600 HH (tablero de KPI)."),
    ("LA-02", "Fase 1", "Procura", "Materiales", "Negativa", "Las órdenes de compra no tenían código de CWP; el almacén no podía reservar materiales por IWP.", "Codificación AWP no incluida en el sistema de compras.", "Alto", "El código de CWP debe estar en toda orden de compra y guía de despacho.", "Campo obligatorio de CWP en el sistema de compras desde la Fase 2.", "Procedimiento de procura", "Fase 2", "Líder de Procura", D(2027, 11, 30), "Implementada", "Pendiente de auditoría al 20 % de la Fase 2."),
    ("LA-03", "Fase 1", "Construcción", "Restricciones", "Negativa", "Los permisos de excavación fueron la restricción más frecuente y se identificaban tarde.", "El permiso no estaba en la lista de restricciones por defecto.", "Alto", "Incluir el permiso como restricción por defecto y solicitarlo 3 semanas antes.", "Actualizar el procedimiento de restricciones y la plantilla de IWP.", "Procedimiento_Gestion_Restricciones.docx", "Fase 2", "Líder de HSE", D(2027, 12, 15), "Verificada", "Sin restricciones de permiso vencidas en Fase 2 (registro)."),
    ("LA-04", "Fase 2", "Ingeniería", "Ingeniería", "Negativa", "La documentación del proveedor de tableros llegó después de la emisión del EWP y retrasó los IWP de montaje.", "La documentación del proveedor no tenía fecha propia en el PWP.", "Medio", "Tratar la documentación del proveedor como entregable del PWP con fecha propia.", "Agregar la fecha de documentación del proveedor en los PWP de la Fase 3.", "Seguimiento_Paquetes.xlsx", "Fase 3", "Líder de Procura", D(2028, 9, 30), "Aprobada", ""),
    ("LA-05", "Fase 2", "Construcción", "Interfaces entre fases", "Positiva", "La reunión conjunta de lookahead con la Fase 3 redujo los conflictos de grúas.", "Coordinación semanal de recursos compartidos.", "Medio", "Mantener un lookahead conjunto durante la superposición de fases.", "Ampliar la reunión a los subcontratistas de la Fase 3.", "Plan de implementación (organización)", "Fase 3", "Gerente de Construcción", D(2028, 8, 31), "Implementada", ""),
    ("LA-06", "Fase 2", "Comisionamiento", "Comisionamiento", "Negativa", "Los IWP no indicaban sistema; fue difícil saber qué faltaba para entregar cada sistema.", "El campo sistema no era obligatorio.", "Alto", "Asignar el sistema a cada IWP desde su creación.", "Campo «Sistema» obligatorio en IWP y en el seguimiento de paquetes.", "Plantilla_IWP.docx", "Fase 3", "Líder de comisionamiento", D(2028, 6, 10), "En análisis", ""),
]
CLAVES_LECCION = ["id", "fase", "etapa", "cat", "tipo", "evento", "causa", "impacto", "leccion", "accion", "doc", "fase_ap", "rol", "f_comp", "estado", "verif"]


def registro_lecciones():
    lib = Libro(SALIDA / "Registro_Lecciones_Aprendidas.xlsx", "Registro de lecciones aprendidas", ["Fases", "Etapas", "Roles"])

    def dias(f, c, p):
        e = f'{c("estado")}{f}'
        return (f'=IF(OR({c("id")}{f}="",{c("f_comp")}{f}=""),"",IF(OR({e}="Implementada",{e}="Verificada",{e}="Descartada"),"",'
                f'{c("f_comp")}{f}-{p["corte"]}))')

    def sit(f, c, p):
        e, d = f'{c("estado")}{f}', f'{c("dias")}{f}'
        return (f'=IF({c("id")}{f}="","",IF(OR({e}="Implementada",{e}="Verificada",{e}="Descartada"),"Cerrada",'
                f'IF({d}="","Sin fecha",IF({d}<0,"Vencida",IF({d}<=15,"Por vencer","En plazo")))))')

    cols = [
        COL_EJEMPLO,
        Col("id", "ID", 7, desc="Identificador (LA-01, LA-02…)."),
        Col("fase", "Fase de origen", 10, lista="Fases", desc="Fase del proyecto en que se produjo la lección."),
        Col("etapa", "Etapa del ciclo de vida", 16, lista="Etapas", desc="Etapa del ciclo de vida relacionada."),
        Col("cat", "Categoría", 17, lista="CategoriasLeccion", desc="Tema de la lección."),
        Col("tipo", "Tipo", 9, lista_valores=["Positiva", "Negativa"], desc="Positiva (práctica a repetir) o Negativa (problema a evitar)."),
        Col("evento", "Evento o situación", 36, ajustar=True, desc="Qué pasó, con datos concretos."),
        Col("causa", "Causa raíz", 26, ajustar=True, desc="Por qué pasó."),
        Col("impacto", "Impacto", 8, lista_valores=dp.IMPACTO_CUALITATIVO, desc="Alto, Medio o Bajo."),
        Col("leccion", "Lección / recomendación", 32, ajustar=True, desc="Qué debe hacerse (o repetirse) en adelante."),
        Col("accion", "Acción propuesta", 32, ajustar=True, desc="Cambio concreto: procedimiento, plantilla, meta, contrato, formación."),
        Col("doc", "Documento afectado", 20, ajustar=True, desc="Documento o plantilla que se actualiza."),
        Col("fase_ap", "Fase de aplicación", 10, lista="Fases", desc="Fase del proyecto en que se aplicará la acción."),
        Col("rol", "Responsable (rol)", 20, lista="Roles", desc="Rol responsable de implementar la acción."),
        Col("f_comp", "Fecha compromiso", 11, "fecha", desc="Fecha límite de implementación (antes de la puerta de control de la fase siguiente)."),
        Col("estado", "Estado", 12, lista="EstadosLeccion", desc="Registrada, En análisis, Aprobada, Implementada, Verificada o Descartada."),
        Col("verif", "Verificación (KPI / evidencia)", 26, ajustar=True, desc="Cómo se comprobó que la acción funcionó."),
        Col("dias", "Días para vencer", 9, "num", formula=dias, desc="Calculado: días hasta la fecha compromiso (negativo si venció). Vacío si está cerrada."),
        Col("sit", "Situación", 11, formula=sit, desc="Calculado: Vencida, Por vencer (≤ 15 días), En plazo o Cerrada."),
    ]
    ws, c, ult = hoja_datos(lib, "Lecciones", "Registro de lecciones aprendidas", cols, ejemplos(LECCIONES, CLAVES_LECCION),
                            n_filas=300, params=[("corte", "Fecha de corte:", dp.FECHA_CORTE, "fecha")], congelar_col=2)
    s = f"${c('sit')}{FILA_DATOS}"
    todo = f"A{FILA_DATOS}:{c('sit')}{ult}"
    resaltar_fila(ws, todo, f'{s}="Vencida"', ROJO, ROJO_TEXTO)
    resaltar_fila(ws, todo, f'{s}="Por vencer"', AMBAR)
    resaltar_fila(ws, f"{c('sit')}{FILA_DATOS}:{c('sit')}{ult}", f'{s}="Cerrada"', VERDE, VERDE_TEXTO)

    rs = lib.wb.create_sheet("Resumen")
    estilo_titulo(rs, "Resumen de lecciones aprendidas", "Calculado a partir de la hoja «Lecciones».", 8)
    rf, re_, rfa = (f"Lecciones!{rango(c, k, ult)}" for k in ("fase", "estado", "fase_ap"))
    tabla_resumen(rs, 4, 1, "Por fase de origen y estado", dp.FASES,
                  [(e, (lambda ref, e=e: f'=COUNTIFS({rf},{ref},{re_},"{e}")'), "0") for e in dp.ESTADOS_LECCION] +
                  [("Vencidas", lambda ref: f'=COUNTIFS({rf},{ref},Lecciones!{rango(c, "sit", ult)},"Vencida")', "0")], 14)
    tabla_resumen(rs, 10, 1, "Por fase de aplicación (acciones que recibe cada fase)", dp.FASES,
                  [("Total", lambda ref: f'=COUNTIF({rfa},{ref})', "0"),
                   ("Pendientes", lambda ref: f'=COUNTIFS({rfa},{ref},{re_},"Registrada")+COUNTIFS({rfa},{ref},{re_},"En análisis")+COUNTIFS({rfa},{ref},{re_},"Aprobada")', "0"),
                   ("Implementadas o verificadas", lambda ref: f'=COUNTIFS({rfa},{ref},{re_},"Implementada")+COUNTIFS({rfa},{ref},{re_},"Verificada")', "0")], 14)
    for i in range(2, 10):
        rs.column_dimensions[get_column_letter(i)].width = 13
    configurar_impresion(rs, "Resumen de lecciones", area="A1:I15")

    hoja_instrucciones(
        lib,
        "Capturar las lecciones aprendidas de cada fase del proyecto y asegurar que se conviertan en acciones implementadas y verificadas en la fase siguiente.",
        "De forma continua; se analiza en la revisión trimestral, en el Comité AWP y al cierre de cada fase (antes de las puertas de control de M14 y M20).",
        "Cualquier miembro del equipo registra; el Líder de WFP consolida; el AWP Champion analiza y el Comité AWP aprueba las acciones.",
        [
            "Actualice la fecha de corte (B3).",
            "Borre las filas amarillas de ejemplo.",
            "Registre el evento con datos concretos y su causa raíz; no culpe a personas.",
            "Formule la lección como recomendación y proponga una acción concreta sobre un documento.",
            "Asigne responsable, fase de aplicación y fecha compromiso.",
            "Actualice el estado hasta «Verificada», con la evidencia (KPI, auditoría).",
        ],
        [("Lecciones", cols)],
        justificacion="El curso 2023 dedica una lección a las lecciones aprendidas y el plan exige transferirlas entre fases superpuestas; el registro las convierte en acciones con dueño y fecha.",
    )
    lib.crear_listas(extra={"CategoriasLeccion": dp.CATEGORIAS_LECCION, "EstadosLeccion": dp.ESTADOS_LECCION})
    lib.guardar()


def main():
    SALIDA.mkdir(parents=True, exist_ok=True)
    for f in (registro_restricciones, matriz_raci, seguimiento_paquetes, definicion_cwa, path_of_construction,
              programa_liberacion, lookahead, checklist_liberacion, tablero_kpi, registro_riesgos, registro_lecciones):
        f()
        print("Generado:", f.__name__)


if __name__ == "__main__":
    main()
